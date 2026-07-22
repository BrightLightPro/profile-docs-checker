"""Core validation logic for the PDF/Excel document table checker.

The checker intentionally does NOT depend on PDF label text. It extracts values by
row position from the fixed table template on a configured PDF page.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, date
from pathlib import Path
from copy import deepcopy
import re
import unicodedata
import json
import contextlib
import io
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import fitz  # PyMuPDF
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

# Internal field names used throughout the program.
PDF_FIELDS = [
    "sprache",
    "dokumentnummer",
    "dok_id",
    "freigabe",
    "artikelnummer",
    "revision",
    "version",
    "ausgabe",
]

PDF_FIELD_DISPLAY = {
    "sprache": "Sprache",
    "dokumentnummer": "Dokumentnummer",
    "dok_id": "DOK-ID",
    "freigabe": "Freigabe-/Änd.-Nr.",
    "artikelnummer": "Artikelnummer",
    "revision": "Revision",
    "version": "Version",
    "ausgabe": "Ausgabe",
    "seiten": "Blatt / sheets",
}

# The row order inside the fixed PDF table. Row 0 = Dokumentname, ignored.
PDF_ROW_INDEX = {
    "sprache": 1,
    "dokumentnummer": 2,
    "dok_id": 3,
    "freigabe": 4,
    "artikelnummer": 5,
    "revision": 6,
    "version": 7,
    "ausgabe": 8,
}

EXCEL_FIELD_TO_HEADER = {
    "sprache": "Sprache",
    "dokumentnummer": "Dok.-Nr.",
    "dok_id": "DOK-ID",
    "freigabe": "Freigabe-/ Änd.-Nr.",
    "artikelnummer": "Artikel-Nr.",
    "revision": "Rev.",
    "version": "Vers.",
}

MATCH_WEIGHTS = {
    "dokumentnummer": 40,
    "artikelnummer": 30,
    "freigabe": 20,
    "revision": 10,
    "version": 10,
    "sprache": 10,
}

MONTH_NAMES = {
    # German
    "januar": 1, "jan": 1,
    "februar": 2, "feb": 2,
    "maerz": 3, "märz": 3, "mar": 3, "mrz": 3,
    "april": 4, "apr": 4,
    "mai": 5,
    "juni": 6, "jun": 6,
    "juli": 7, "jul": 7,
    "august": 8, "aug": 8,
    "september": 9, "sep": 9, "sept": 9,
    "oktober": 10, "okt": 10, "oct": 10,
    "november": 11, "nov": 11,
    "dezember": 12, "dez": 12, "dec": 12,
    # English
    "january": 1,
    "february": 2,
    "march": 3,
    "may": 5,
    "june": 6,
    "july": 7,
    "october": 10,
    "december": 12,
    # French
    "janvier": 1, "fevrier": 2, "février": 2, "mars": 3, "avril": 4,
    "juin": 6, "juillet": 7, "aout": 8, "août": 8,
    "septembre": 9, "octobre": 10, "novembre": 11, "decembre": 12, "décembre": 12,
    # Spanish / Portuguese / Italian common month names
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5,
    "junio": 6, "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9,
    "octubre": 10, "noviembre": 11, "diciembre": 12,
    "janeiro": 1, "fevereiro": 2, "marco": 3, "março": 3, "maio": 5,
    "julho": 7, "setembro": 9, "outubro": 10, "dezembro": 12,
    "gennaio": 1, "febbraio": 2, "maggio": 5, "giugno": 6,
    "luglio": 7, "ottobre": 10, "dicembre": 12,
}

STATUS_ORDER = {
    "OK": 1,
    "OK_WITH_NORMALIZATION": 2,
    "AUSGABE_MISMATCH": 3,
    "MISMATCH": 4,
    "MISMATCH_AND_AUSGABE_MISMATCH": 5,
    "LIKELY_WRONG_DOK_ID": 6,
    "DUPLICATE_DOK_ID_RESOLVED": 7,
    "DUPLICATE_DOK_ID_AMBIGUOUS": 8,
    "DOK_ID_NOT_FOUND": 9,
    "NO_RELIABLE_MATCH": 10,
    "EXTRACTION_ERROR": 11,
}


@dataclass
class ExtractedPDF:
    file_path: Path
    file_name: str
    values_raw: Dict[str, str] = field(default_factory=dict)
    values_norm: Dict[str, str] = field(default_factory=dict)
    page_number: int = 2  # Human visible page number, 1-based.
    table_index: Optional[int] = None  # 1-based within detected tables.
    row_start: Optional[int] = None  # 1-based row where Dokumentname starts within detected table.
    error: Optional[str] = None
    warnings: List[str] = field(default_factory=list)
    total_pages: Optional[int] = None


@dataclass
class ExcelRecord:
    row_number: int
    values_raw: Dict[str, str]
    values_norm: Dict[str, str]


@dataclass
class FieldComparison:
    file_name: str
    field: str
    source: str  # "Excel" or "User input"
    pdf_raw: str
    pdf_norm: str
    reference_raw: str
    reference_norm: str
    result: str
    excel_row: Optional[int] = None


@dataclass
class ValidationResult:
    extracted: ExtractedPDF
    status: str
    confidence: float = 0.0
    matched_excel_row: Optional[int] = None
    issues: List[str] = field(default_factory=list)
    comparisons: List[FieldComparison] = field(default_factory=list)
    candidate_rows: List[Tuple[int, float]] = field(default_factory=list)


@dataclass
class WordMapping:
    table_index_zero_based: int
    header_row_zero_based: int
    data_start_row_zero_based: int
    columns_by_index_zero_based: Dict[str, int]


@dataclass
class WordExtractionInfo:
    mode: str = "custom_mapping"
    template_type: str = "Custom mapping"
    table_index_zero_based: int = 0
    header_rows_zero_based: List[int] = field(default_factory=list)
    data_start_row_zero_based: int = 0
    columns_by_index_zero_based: Dict[str, int] = field(default_factory=dict)
    change_number_raw: str = ""
    change_number_norm: str = ""
    change_number_source: str = "body_table"
    change_number_section_zero_based: Optional[int] = None
    change_number_header_table_zero_based: Optional[int] = None
    change_number_cell_row_zero_based: Optional[int] = None
    change_number_cell_col_zero_based: Optional[int] = None
    warnings: List[str] = field(default_factory=list)


@dataclass
class WordRecord:
    word_file_name: str
    word_row_number: int  # 1-based Word table row number for human review.
    values_raw: Dict[str, str]
    values_norm: Dict[str, str]


@dataclass
class WordFieldComparison:
    word_file_name: str
    word_row_number: int
    field: str
    word_raw: str
    word_norm: str
    excel_raw: str
    excel_norm: str
    result: str
    excel_row: Optional[int] = None


@dataclass
class WordValidationResult:
    record: WordRecord
    status: str
    confidence: float = 0.0
    matched_excel_row: Optional[int] = None
    issues: List[str] = field(default_factory=list)
    comparisons: List[WordFieldComparison] = field(default_factory=list)
    candidate_rows: List[Tuple[int, float]] = field(default_factory=list)
    excel_status: str = ""
    page_count_result: str = "NOT_CHECKED"
    word_pages_raw: str = ""
    word_pages_norm: str = ""
    pdf_pages: Optional[int] = None
    pdf_file_name: str = ""
    page_count_issue: str = ""


@dataclass
class WordCorrectionResult:
    output_path: Path
    changed_cells: int = 0
    corrected_rows: int = 0
    skipped_rows: int = 0
    warnings: List[str] = field(default_factory=list)


@dataclass
class WordFillResult:
    output_path: Path
    filled_rows: int = 0
    added_rows: int = 0
    cleared_rows: int = 0
    changed_cells: int = 0
    warnings: List[str] = field(default_factory=list)


WORD_STATUS_ORDER = {
    "WORD_OK": 1,
    "WORD_OK_WITH_NORMALIZATION": 2,
    "WORD_DUPLICATE_DOCUMENT_NUMBER_RESOLVED": 3,
    "WORD_LIKELY_WRONG_DOCUMENT_NUMBER": 4,
    "WORD_LIKELY_WRONG_DOK_ID": 5,
    "WORD_CHANGE_NUMBER_MISMATCH": 6,
    "WORD_FIELD_MISMATCH": 6,
    "WORD_FIELD_AND_CHANGE_NUMBER_MISMATCH": 7,
    "WORD_DOCUMENT_NOT_IN_EXCEL": 8,
    "WORD_AMBIGUOUS_MATCH": 9,
    "WORD_NO_RELIABLE_MATCH": 10,
    "WORD_TABLE_NOT_RECOGNIZED": 11,
    "WORD_PAGE_COUNT_MISMATCH": 7,
    "WORD_MULTIPLE_MISMATCHES": 8,
    "WORD_EXTRACTION_ERROR": 12,
}



def clean_text(value: object) -> str:
    """Return a stable one-line string without changing semantic content too much."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    text = str(value)
    # Preserve meaningful characters exactly. Only remove layout artifacts.
    text = text.replace("\u00a0", " ")
    text = text.replace("\r", " ").replace("\n", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_identifier(value: object) -> str:
    """Preserve the document value exactly, apart from technical whitespace cleanup.

    Identifiers, document numbers, article numbers, change numbers, revisions and
    versions are intentionally compared case-sensitively and character-for-character.
    """
    return clean_text(value)


def normalize_header(value: object) -> str:
    # Header detection may be tolerant; document values are not.
    text = unicodedata.normalize("NFKC", clean_text(value)).casefold()
    text = text.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    return re.sub(r"[^a-z0-9]+", "", text)


def normalize_language(value: object) -> str:
    """Map PDF language_country codes to the two-letter codes used by the Excel list.

    The local system uses ``jp`` for Japanese, while ISO language tags commonly use
    ``ja``. Both ``ja_JP`` and ``jp_JP`` therefore compare to Excel ``jp``.
    """
    text = clean_text(value).casefold().replace("-", "_")
    text_no_space = re.sub(r"\s+", "", text)
    m = re.match(r"^([a-z]{2})(?:_[a-z]{2})?$", text_no_space)
    if m:
        code = m.group(1)
    else:
        pairs = re.findall(r"[a-z]{2}", text)
        code = pairs[0] if pairs else text_no_space
    return "jp" if code in {"ja", "jp"} else code


def normalize_revision_version(value: object) -> str:
    # Revisions and versions are strict values: 01 is not the same as 1.
    return clean_text(value)


def normalize_page_count(value: object) -> str:
    """Normalize a Word sheet/page count to a positive integer string."""
    text = clean_text(value)
    if not text:
        return ""
    # Common forms: 3, 03, 3 pages, 3 Blatt.
    m = re.fullmatch(r"0*(\d+)(?:\s*(?:pages?|sheets?|blatt|seiten?))?", text, flags=re.I)
    if m:
        number = int(m.group(1))
    else:
        # Also accept a sheet position that includes a total, e.g. 1/3 or 1 of 3.
        total = re.fullmatch(r"0*\d+\s*(?:/|of|von)\s*0*(\d+)", text, flags=re.I)
        if not total:
            return ""
        number = int(total.group(1))
    return str(number) if number > 0 else ""


def _normalize_month_number(month: int, year: int) -> Optional[str]:
    if 1 <= month <= 12 and 1900 <= year <= 2100:
        return f"{year:04d}-{month:02d}"
    return None


def normalize_month_year(value: object, allow_two_digit_year: bool = True) -> str:
    """Normalize many month/year formats to YYYY-MM, or return empty string if unsupported."""
    text = clean_text(value).casefold()
    if not text:
        return ""
    text = text.replace("–", "-").replace("—", "-").replace("−", "-")
    text = re.sub(r"\s+", " ", text).strip()

    # Date objects from Excel/openpyxl etc.
    if isinstance(value, (datetime, date)):
        return f"{value.year:04d}-{value.month:02d}"

    # Month-first, preferred: MM.YYYY / MM/YYYY / MM-YYYY / MM YYYY
    m = re.fullmatch(r"(\d{1,2})[.\-/\s](\d{4})", text)
    if m:
        return _normalize_month_number(int(m.group(1)), int(m.group(2))) or ""

    # Year-first fallback: YYYY-MM / YYYY/MM / YYYY.MM / YYYY MM
    m = re.fullmatch(r"(\d{4})[.\-/\s](\d{1,2})", text)
    if m:
        return _normalize_month_number(int(m.group(2)), int(m.group(1))) or ""

    if allow_two_digit_year:
        # Two-digit years are treated as 2000-2099. This is conservative for current technical docs.
        m = re.fullmatch(r"(\d{1,2})[.\-/\s](\d{2})", text)
        if m:
            return _normalize_month_number(int(m.group(1)), 2000 + int(m.group(2))) or ""

    # Month name + year, or year + month name.
    m = re.fullmatch(r"([a-zäöüßéèêûùàáíóúçãõ]+)\.?\s+(\d{4})", text)
    if m:
        month = MONTH_NAMES.get(m.group(1))
        if month:
            return _normalize_month_number(month, int(m.group(2))) or ""
    m = re.fullmatch(r"(\d{4})\s+([a-zäöüßéèêûùàáíóúçãõ]+)\.?,?", text)
    if m:
        month = MONTH_NAMES.get(m.group(2))
        if month:
            return _normalize_month_number(month, int(m.group(1))) or ""

    return ""


def normalize_expected_ausgabe(user_value: str) -> str:
    """Expected input: month + year, normalized to YYYY-MM.

    Preferred user format is MM.YYYY, but MM-YYYY, MM/YYYY and MM YYYY
    are accepted to avoid unnecessary input failures.
    """
    raw = clean_text(user_value)
    m = re.fullmatch(r"(\d{1,2})[.\-/\s](\d{4})", raw)
    if not m:
        raise ValueError("Expected Ausgabe must be entered as month and year, for example 07.2026, 07-2026 or 07/2026")
    normalized = _normalize_month_number(int(m.group(1)), int(m.group(2)))
    if not normalized:
        raise ValueError("Expected Ausgabe has an invalid month or year. Use for example 07.2026")
    return normalized


def normalize_field(field_name: str, value: object) -> str:
    # Only language, Ausgabe and page count have defined semantic conversion rules.
    # Every other field is compared exactly after technical whitespace cleanup.
    if field_name == "sprache":
        return normalize_language(value)
    if field_name == "ausgabe":
        return normalize_month_year(value)
    if field_name == "seiten":
        return normalize_page_count(value)
    return clean_text(value)


def excel_cell_to_text(cell) -> str:
    """Read Excel cell text while preserving obvious zero-padded numeric formats."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        text = str(value)
        # Preserve simple zero-padded number formats like 000000 or 0000000.
        number_format = str(cell.number_format or "")
        if re.fullmatch(r"0+", number_format) and isinstance(value, int):
            text = str(value).zfill(len(number_format))
        return text
    return clean_text(value)


def load_excel_records(excel_path: Path | str, sheet_name: Optional[str] = None) -> Tuple[List[ExcelRecord], str, int, Dict[str, int]]:
    """Load relevant Excel columns from the fixed German template.

    Returns: records, actual_sheet_name, header_row_number, columns_by_field.
    """
    excel_path = Path(excel_path)
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    wb = load_workbook(excel_path, data_only=True, read_only=False)
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found. Available sheets: {', '.join(wb.sheetnames)}")
        ws = wb[sheet_name]
    else:
        ws = wb[wb.sheetnames[0]]

    header_variants = {normalize_header(header): field for field, header in EXCEL_FIELD_TO_HEADER.items()}
    # A few tolerant variants for the fixed template.
    header_variants.update({
        normalize_header("Dok Nr"): "dokumentnummer",
        normalize_header("Dokumentnummer"): "dokumentnummer",
        normalize_header("Dokument-Nr."): "dokumentnummer",
        normalize_header("Dok-ID"): "dok_id",
        normalize_header("DOK ID"): "dok_id",
        normalize_header("Freigabe-/Änd.-Nr."): "freigabe",
        normalize_header("Freigabe / Änd. Nr."): "freigabe",
        normalize_header("Artikel Nr."): "artikelnummer",
        normalize_header("Artikelnummer"): "artikelnummer",
        normalize_header("Revision"): "revision",
        normalize_header("Version"): "version",
    })

    best_header_row = None
    best_columns: Dict[str, int] = {}
    best_count = 0
    max_scan_row = min(ws.max_row, 50)
    for row in range(1, max_scan_row + 1):
        columns: Dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            key = normalize_header(ws.cell(row, col).value)
            if key in header_variants:
                columns[header_variants[key]] = col
        if len(columns) > best_count:
            best_count = len(columns)
            best_header_row = row
            best_columns = columns
        if best_count == len(EXCEL_FIELD_TO_HEADER):
            break

    missing = [EXCEL_FIELD_TO_HEADER[f] for f in EXCEL_FIELD_TO_HEADER if f not in best_columns]
    if best_header_row is None or missing:
        raise ValueError(
            "Could not find all required Excel columns. Missing: "
            + ", ".join(missing)
            + ". Check the sheet/template."
        )

    records: List[ExcelRecord] = []
    for row in range(best_header_row + 1, ws.max_row + 1):
        raw: Dict[str, str] = {}
        norm: Dict[str, str] = {}
        for field_name, col in best_columns.items():
            raw_value = excel_cell_to_text(ws.cell(row, col))
            raw[field_name] = raw_value
            norm[field_name] = normalize_field(field_name, raw_value)
        if any(raw.get(f, "") for f in EXCEL_FIELD_TO_HEADER):
            records.append(ExcelRecord(row_number=row, values_raw=raw, values_norm=norm))

    return records, ws.title, best_header_row, best_columns


def _last_nonempty_cell_after_label(row: Sequence[object]) -> str:
    """For a table row, take the last non-empty cell after the label column."""
    cells = [clean_text(c) for c in row]
    if len(cells) <= 1:
        return cells[0] if cells else ""
    right_side = [c for c in cells[1:] if c]
    if right_side:
        return right_side[-1]
    # If table extraction merged the row unexpectedly, return last non-empty cell.
    non_empty = [c for c in cells if c]
    return non_empty[-1] if non_empty else ""


def _candidate_score_from_rows(rows: Sequence[Sequence[object]], start: int) -> float:
    """Score a possible 9-row metadata table slice. Higher is better."""
    if start < 0 or start + 8 >= len(rows):
        return -1
    raw_values = {field_name: _last_nonempty_cell_after_label(rows[start + idx]) for field_name, idx in PDF_ROW_INDEX.items()}
    score = 0.0
    # Positive signals without relying on labels.
    if re.fullmatch(r"[A-Za-z]{2}[_-][A-Za-z]{2}", clean_text(raw_values.get("sprache", ""))):
        score += 35
    elif raw_values.get("sprache"):
        score += 10
    if raw_values.get("dok_id"):
        score += 25
    if raw_values.get("dokumentnummer"):
        score += 15
    if raw_values.get("artikelnummer"):
        score += 15
    if normalize_month_year(raw_values.get("ausgabe", "")):
        score += 15
    # Prefer slices with values in most relevant rows.
    score += sum(2 for v in raw_values.values() if clean_text(v))
    return score


def _extract_from_table_rows(
    rows: Sequence[Sequence[object]],
    file_path: Path,
    page_number: int,
    table_index: int,
    forced_row_start: Optional[int] = None,
) -> ExtractedPDF:
    if forced_row_start is not None:
        start = forced_row_start - 1
    else:
        possible = [(i, _candidate_score_from_rows(rows, i)) for i in range(0, max(0, len(rows) - 8))]
        if not possible:
            raise ValueError("Detected table has fewer than 9 rows.")
        start, _score = max(possible, key=lambda x: x[1])
    if start < 0 or start + 8 >= len(rows):
        raise ValueError(f"Row start {start + 1} is invalid for a table with {len(rows)} rows.")

    raw: Dict[str, str] = {}
    norm: Dict[str, str] = {}
    for field_name, idx in PDF_ROW_INDEX.items():
        value = _last_nonempty_cell_after_label(rows[start + idx])
        raw[field_name] = value
        norm[field_name] = normalize_field(field_name, value)

    warnings: List[str] = []
    if not norm.get("sprache"):
        warnings.append("Could not normalize Sprache from PDF.")
    if not raw.get("dok_id"):
        warnings.append("DOK-ID is empty in PDF extraction.")
    if raw.get("ausgabe") and not norm.get("ausgabe"):
        warnings.append("Could not normalize Ausgabe from PDF.")

    return ExtractedPDF(
        file_path=file_path,
        file_name=file_path.name,
        values_raw=raw,
        values_norm=norm,
        page_number=page_number,
        table_index=table_index,
        row_start=start + 1,
        warnings=warnings,
    )


def extract_pdf_table(
    pdf_path: Path | str,
    page_number: int = 2,
    table_index: Optional[int] = None,
    row_start: Optional[int] = None,
) -> ExtractedPDF:
    """Extract values from the fixed PDF table on a given visible page number.

    table_index and row_start are 1-based optional overrides. Leave them blank for auto detection.
    """
    pdf_path = Path(pdf_path)
    result = ExtractedPDF(file_path=pdf_path, file_name=pdf_path.name, page_number=page_number)
    try:
        if not pdf_path.exists():
            raise FileNotFoundError(str(pdf_path))
        doc = fitz.open(pdf_path)
        result.total_pages = len(doc)
        try:
            if page_number < 1 or page_number > len(doc):
                raise ValueError(f"PDF has {len(doc)} pages; requested page {page_number}.")
            page = doc[page_number - 1]
            # PyMuPDF may print an optional package recommendation during table detection.
            # Suppress it so CLI/GUI output stays focused on validation results.
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                found = page.find_tables()
            tables = list(found.tables or [])
            if not tables:
                raise ValueError("No table detected on the configured page. Is the page number correct?")

            candidates: List[Tuple[float, int, Sequence[Sequence[object]]]] = []
            selected_tables = []
            if table_index is not None:
                if table_index < 1 or table_index > len(tables):
                    raise ValueError(f"Table index {table_index} invalid. Detected {len(tables)} table(s) on the page.")
                selected_tables = [(table_index, tables[table_index - 1])]
            else:
                selected_tables = [(i + 1, t) for i, t in enumerate(tables)]

            for idx, table in selected_tables:
                rows = table.extract()
                if not rows:
                    continue
                if row_start is not None:
                    score = _candidate_score_from_rows(rows, row_start - 1)
                    candidates.append((score, idx, rows))
                else:
                    max_score = max((_candidate_score_from_rows(rows, s) for s in range(0, max(0, len(rows) - 8))), default=-1)
                    # Slightly prefer larger tables because the metadata table has many rows.
                    max_score += min(len(rows), 20) * 0.1
                    candidates.append((max_score, idx, rows))

            if not candidates:
                raise ValueError("No readable table rows found on the configured page.")
            _score, chosen_idx, rows = max(candidates, key=lambda x: x[0])
            extracted = _extract_from_table_rows(rows, pdf_path, page_number, chosen_idx, row_start)
            extracted.total_pages = len(doc)
            return extracted
        finally:
            doc.close()
    except Exception as exc:
        result.error = str(exc)
        return result


def compare_value(field_name: str, pdf_raw: str, reference_raw: str, source: str, file_name: str, excel_row: Optional[int]) -> FieldComparison:
    pdf_clean = clean_text(pdf_raw)
    ref_clean = clean_text(reference_raw)
    pdf_norm = normalize_field(field_name, pdf_raw)
    ref_norm = normalize_field(field_name, reference_raw)
    if field_name == "ausgabe":
        ref_norm = normalize_month_year(reference_raw)

    if not pdf_clean and not ref_clean:
        res = "BOTH_EMPTY"
    elif not pdf_clean:
        res = "PDF_MISSING"
    elif not ref_clean:
        res = "REFERENCE_MISSING"
    elif field_name == "ausgabe" and not pdf_norm:
        res = "PDF_UNSUPPORTED_DATE_FORMAT"
    elif field_name in {"sprache", "ausgabe"}:
        res = "OK" if pdf_norm and pdf_norm == ref_norm else "MISMATCH"
    else:
        res = "OK" if pdf_clean == ref_clean else "MISMATCH"
    return FieldComparison(
        file_name=file_name,
        field=PDF_FIELD_DISPLAY[field_name],
        source=source,
        pdf_raw=pdf_clean,
        pdf_norm=pdf_norm,
        reference_raw=ref_clean,
        reference_norm=ref_norm,
        result=res,
        excel_row=excel_row,
    )


def score_candidate(pdf: ExtractedPDF, record: ExcelRecord) -> float:
    earned = 0
    possible = 0
    for field_name, weight in MATCH_WEIGHTS.items():
        pdf_norm = pdf.values_norm.get(field_name, "")
        xl_norm = record.values_norm.get(field_name, "")
        if not pdf_norm or not xl_norm:
            continue
        possible += weight
        if pdf_norm == xl_norm:
            earned += weight
    if possible == 0:
        return 0.0
    return round(100.0 * earned / possible, 1)


def build_dok_id_index(records: Iterable[ExcelRecord]) -> Dict[str, List[ExcelRecord]]:
    idx: Dict[str, List[ExcelRecord]] = {}
    for rec in records:
        dok = rec.values_norm.get("dok_id", "")
        if dok:
            idx.setdefault(dok, []).append(rec)
    return idx


def _candidate_list(pdf: ExtractedPDF, records: Sequence[ExcelRecord]) -> List[Tuple[ExcelRecord, float]]:
    scored = [(rec, score_candidate(pdf, rec)) for rec in records]
    scored.sort(key=lambda x: x[1], reverse=True)
    return scored


def _best_candidate_is_clear(scored: List[Tuple[ExcelRecord, float]], threshold: float = 70.0, margin: float = 15.0) -> bool:
    if not scored:
        return False
    best = scored[0][1]
    second = scored[1][1] if len(scored) > 1 else 0
    return best >= threshold and (best - second >= margin or best >= 90.0)


def validate_one_pdf(
    extracted: ExtractedPDF,
    records: Sequence[ExcelRecord],
    dok_id_index: Dict[str, List[ExcelRecord]],
    expected_ausgabe_input: str,
    expected_ausgabe_norm: str,
) -> ValidationResult:
    if extracted.error:
        return ValidationResult(
            extracted=extracted,
            status="EXTRACTION_ERROR",
            issues=[extracted.error],
        )

    result = ValidationResult(extracted=extracted, status="NO_RELIABLE_MATCH")

    # Ausgabe is always checked against user input, independently of Excel.
    ausgabe_comp = FieldComparison(
        file_name=extracted.file_name,
        field=PDF_FIELD_DISPLAY["ausgabe"],
        source="User input",
        pdf_raw=extracted.values_raw.get("ausgabe", ""),
        pdf_norm=extracted.values_norm.get("ausgabe", ""),
        reference_raw=expected_ausgabe_input,
        reference_norm=expected_ausgabe_norm,
        result=(
            "OK" if extracted.values_norm.get("ausgabe", "") == expected_ausgabe_norm
            else (
                "PDF_UNSUPPORTED_DATE_FORMAT"
                if extracted.values_raw.get("ausgabe") and not extracted.values_norm.get("ausgabe")
                else "MISMATCH"
            )
        ),
    )

    pdf_dok = extracted.values_norm.get("dok_id", "")
    matched_record: Optional[ExcelRecord] = None
    match_reason = ""

    if pdf_dok and pdf_dok in dok_id_index:
        rows = dok_id_index[pdf_dok]
        if len(rows) == 1:
            matched_record = rows[0]
            result.confidence = 100.0
            match_reason = "DOK-ID found exactly once in Excel."
        else:
            scored = _candidate_list(extracted, rows)
            result.candidate_rows = [(rec.row_number, score) for rec, score in scored[:10]]
            if _best_candidate_is_clear(scored, threshold=70.0, margin=10.0):
                matched_record = scored[0][0]
                result.confidence = scored[0][1]
                result.status = "DUPLICATE_DOK_ID_RESOLVED"
                match_reason = f"DOK-ID appears {len(rows)} times; best duplicate row selected."
                result.issues.append(match_reason)
            else:
                result.status = "DUPLICATE_DOK_ID_AMBIGUOUS"
                result.confidence = scored[0][1] if scored else 0.0
                result.issues.append(f"DOK-ID appears {len(rows)} times and no unique best row was found.")
                result.comparisons.append(ausgabe_comp)
                return result
    else:
        scored = _candidate_list(extracted, records)
        result.candidate_rows = [(rec.row_number, score) for rec, score in scored[:10]]
        if _best_candidate_is_clear(scored, threshold=70.0, margin=15.0):
            matched_record = scored[0][0]
            result.confidence = scored[0][1]
            result.status = "LIKELY_WRONG_DOK_ID"
            issue = "DOK-ID from PDF is empty or not present in Excel, but another row strongly matches other fields."
            result.issues.append(issue)
            match_reason = issue
        else:
            result.confidence = scored[0][1] if scored else 0.0
            result.status = "DOK_ID_NOT_FOUND" if pdf_dok else "NO_RELIABLE_MATCH"
            if pdf_dok:
                result.issues.append("DOK-ID from PDF was not found in Excel and no reliable fallback match was found.")
            else:
                result.issues.append("DOK-ID in PDF is empty and no reliable fallback match was found.")
            result.comparisons.append(ausgabe_comp)
            return result

    result.matched_excel_row = matched_record.row_number if matched_record else None

    # Compare PDF fields with Excel row.
    comps: List[FieldComparison] = []
    if matched_record:
        for field_name in ["sprache", "dokumentnummer", "dok_id", "freigabe", "artikelnummer", "revision", "version"]:
            comp = compare_value(
                field_name,
                extracted.values_raw.get(field_name, ""),
                matched_record.values_raw.get(field_name, ""),
                source="Excel",
                file_name=extracted.file_name,
                excel_row=matched_record.row_number,
            )
            comps.append(comp)
    comps.append(ausgabe_comp)
    result.comparisons = comps

    excel_bad = [c for c in comps if c.source == "Excel" and c.result not in {"OK", "BOTH_EMPTY"}]
    ausgabe_bad = ausgabe_comp.result != "OK"

    if result.status == "DUPLICATE_DOK_ID_RESOLVED":
        if excel_bad:
            result.issues.append("Best duplicate row still has mismatching Excel fields: " + ", ".join(c.field for c in excel_bad))
        if ausgabe_bad:
            result.issues.append("Ausgabe differs from expected input.")
        return result

    if result.status == "LIKELY_WRONG_DOK_ID":
        if excel_bad:
            result.issues.append("Fallback row found, but some Excel fields still differ: " + ", ".join(c.field for c in excel_bad))
        if ausgabe_bad:
            result.issues.append("Ausgabe differs from expected input.")
        return result

    if excel_bad and ausgabe_bad:
        result.status = "MISMATCH_AND_AUSGABE_MISMATCH"
        result.issues.append("Excel fields differ and Ausgabe differs from expected input.")
    elif excel_bad:
        result.status = "MISMATCH"
        result.issues.append("Mismatching Excel fields: " + ", ".join(c.field for c in excel_bad))
    elif ausgabe_bad:
        result.status = "AUSGABE_MISMATCH"
        result.issues.append("Ausgabe differs from expected input.")
    else:
        result.status = "OK"
        if match_reason:
            result.issues.append(match_reason)
    return result


def list_pdf_files(pdf_folder: Path | str) -> List[Path]:
    folder = Path(pdf_folder)
    if not folder.exists() or not folder.is_dir():
        raise NotADirectoryError(f"PDF folder not found: {folder}")
    return sorted([p for p in folder.iterdir() if p.is_file() and p.suffix.casefold() == ".pdf"])




# ---------------------------------------------------------------------------
# Word (.docx) table inspection and Word-vs-Excel comparison
# ---------------------------------------------------------------------------

WORD_CUSTOM_COMPARISON_FIELDS = [
    "sprache", "dokumentnummer", "dok_id", "freigabe", "artikelnummer", "revision", "version"
]
WORD_TEMPLATE_COMPARISON_FIELDS = ["dokumentnummer", "artikelnummer", "revision", "version", "freigabe"]
WORD_REPORT_FIELDS = ["dokumentnummer", "artikelnummer", "revision", "version", "freigabe", "seiten"]
WORD_EXCEL_COMPARISON_FIELDS = ["dokumentnummer", "artikelnummer", "revision", "version", "freigabe", "dok_id", "sprache"]
WORD_COMPARISON_FIELDS = WORD_CUSTOM_COMPARISON_FIELDS  # Backward-compatible public name.
WORD_MATCH_WEIGHTS = {
    "dokumentnummer": 60,
    "artikelnummer": 25,
    "revision": 8,
    "version": 7,
}


def _word_cell_to_text(cell) -> str:
    """Read all visible text in a Word cell, including content controls/text boxes.

    ``python-docx``'s ``cell.text`` can miss text nested inside structured document
    tags. The change-notice Number field is often implemented that way, so read
    every ``w:t`` descendant from the cell XML instead.
    """
    try:
        paragraph_texts: List[str] = []
        for paragraph in cell._tc.xpath(".//w:p"):
            # Text split across styled runs belongs together; paragraph boundaries
            # become a single technical space.
            paragraph_text = "".join((node.text or "") for node in paragraph.xpath(".//w:t"))
            if paragraph_text:
                paragraph_texts.append(paragraph_text)
        xml_text = clean_text(" ".join(paragraph_texts))
        if xml_text:
            return xml_text
    except Exception:
        pass
    return clean_text(cell.text)


def _set_word_cell_text_preserve_structure(cell, value: object) -> None:
    """Replace visible text while preserving the existing cell/table formatting.

    Existing text nodes are reused so cell borders, shading, paragraph alignment,
    content controls and run formatting remain in place.
    """
    text = clean_text(value)
    try:
        nodes = list(cell._tc.xpath(".//w:t"))
    except Exception:
        nodes = []
    if nodes:
        nodes[0].text = text
        for node in nodes[1:]:
            node.text = ""
    else:
        cell.text = text


def _field_header_variants() -> Dict[str, str]:
    variants = {normalize_header(header): field for field, header in EXCEL_FIELD_TO_HEADER.items()}
    variants.update({
        normalize_header("Sprache"): "sprache",
        normalize_header("Language"): "sprache",
        normalize_header("Langue"): "sprache",
        normalize_header("Idioma"): "sprache",
        normalize_header("Dok Nr"): "dokumentnummer",
        normalize_header("Dok.-Nr."): "dokumentnummer",
        normalize_header("Dokumentnummer"): "dokumentnummer",
        normalize_header("Dokument-Nr."): "dokumentnummer",
        normalize_header("Document No."): "dokumentnummer",
        normalize_header("Doc No."): "dokumentnummer",
        normalize_header("DOK-ID"): "dok_id",
        normalize_header("Dok-ID"): "dok_id",
        normalize_header("DOK ID"): "dok_id",
        normalize_header("Doc ID"): "dok_id",
        normalize_header("Freigabe-/Änd.-Nr."): "freigabe",
        normalize_header("Freigabe-/ Änd.-Nr."): "freigabe",
        normalize_header("Freigabe / Änd. Nr."): "freigabe",
        normalize_header("Release/Change No."): "freigabe",
        normalize_header("Artikel Nr."): "artikelnummer",
        normalize_header("Artikel-Nr."): "artikelnummer",
        normalize_header("Artikelnummer"): "artikelnummer",
        normalize_header("Article No."): "artikelnummer",
        normalize_header("Part No."): "artikelnummer",
        normalize_header("Revision"): "revision",
        normalize_header("Rev."): "revision",
        normalize_header("Rev"): "revision",
        normalize_header("Version"): "version",
        normalize_header("Vers."): "version",
        normalize_header("Vers"): "version",
    })
    return variants


def _detect_columns_from_header(headers: Sequence[str]) -> Dict[str, int]:
    variants = _field_header_variants()
    found: Dict[str, int] = {}
    for idx, header in enumerate(headers):
        key = normalize_header(header)
        if key in variants and variants[key] not in found:
            found[variants[key]] = idx
    return found


def inspect_word_tables(docx_path: Path | str, include_samples: bool = False, sample_rows: int = 5) -> Dict[str, Any]:
    """Return a privacy-conscious structural description of tables in a DOCX."""
    docx_path = Path(docx_path)
    if not docx_path.exists():
        raise FileNotFoundError(f"Word file not found: {docx_path}")
    doc = Document(docx_path)
    result: Dict[str, Any] = {"file_name": docx_path.name, "tables": []}
    for table_index, table in enumerate(doc.tables):
        rows: List[List[str]] = []
        max_cols = 0
        for row in table.rows:
            cells = [_word_cell_to_text(cell) for cell in row.cells]
            max_cols = max(max_cols, len(cells))
            rows.append(cells)

        best_header_row = 0 if rows else None
        best_count = -1
        best_columns: Dict[str, int] = {}
        for i, cells in enumerate(rows[: min(len(rows), 15)]):
            cols = _detect_columns_from_header(cells)
            if len(cols) > best_count:
                best_count = len(cols)
                best_header_row = i
                best_columns = cols

        auto_detect = _detect_change_notice_table(table, table_index)
        sample_payload = []
        if rows and best_header_row is not None:
            start = best_header_row + 1
            for r in rows[start : start + sample_rows]:
                sample_payload.append(r if include_samples else ["<non-empty>" if clean_text(c) else "" for c in r])

        payload = {
            "table_index_zero_based": table_index,
            "row_count": len(rows),
            "column_count": max_cols,
            "likely_header_row_zero_based": best_header_row,
            "headers": rows[best_header_row] if rows and best_header_row is not None else [],
            "detected_columns_by_index_zero_based": best_columns,
            "suggested_mapping": {
                "table_index_zero_based": table_index,
                "header_row_zero_based": best_header_row if best_header_row is not None else 0,
                "data_start_row_zero_based": (best_header_row + 1) if best_header_row is not None else 1,
                "columns_by_index_zero_based": {field: best_columns.get(field) for field in WORD_CUSTOM_COMPARISON_FIELDS},
            },
            "sample_rows": sample_payload,
        }
        if auto_detect:
            payload["recognized_change_notice_template"] = {
                "score": auto_detect["score"],
                "header_rows_zero_based": auto_detect["header_rows"],
                "data_start_row_zero_based": auto_detect["data_start"],
                "columns_by_index_zero_based": auto_detect["columns"],
                "change_number_raw": auto_detect["change_number_raw"],
                "change_number_cell_row_zero_based": auto_detect.get("change_number_cell_row"),
                "change_number_cell_col_zero_based": auto_detect.get("change_number_cell_col"),
                "warnings": auto_detect["warnings"],
            }
        result["tables"].append(payload)
    return result


def write_word_inspection_files(
    docx_path: Path | str,
    structure_out: Path | str = "word_table_structure.json",
    mapping_out: Path | str = "word_mapping_template.json",
    include_samples: bool = False,
) -> Tuple[Path, Path]:
    structure = inspect_word_tables(docx_path, include_samples=include_samples)
    structure_out = Path(structure_out)
    mapping_out = Path(mapping_out)
    structure_out.parent.mkdir(parents=True, exist_ok=True)
    mapping_out.parent.mkdir(parents=True, exist_ok=True)
    structure_out.write_text(json.dumps(structure, ensure_ascii=False, indent=2), encoding="utf-8")

    tables = structure.get("tables", [])
    if tables:
        chosen = max(tables, key=lambda t: len([v for v in t.get("suggested_mapping", {}).get("columns_by_index_zero_based", {}).values() if v is not None]))
        mapping = chosen.get("suggested_mapping", {})
    else:
        mapping = {
            "table_index_zero_based": 0,
            "header_row_zero_based": 0,
            "data_start_row_zero_based": 1,
            "columns_by_index_zero_based": {field: None for field in WORD_CUSTOM_COMPARISON_FIELDS},
        }
    mapping["_instructions"] = (
        "This mapping is only needed for non-standard Word tables. The Freigabe-/Änderungsmitteilung / "
        "Engineering Change Notice template is recognized automatically without a mapping file. "
        "Indexes are zero-based; leave a field null to ignore it."
    )
    mapping_out.write_text(json.dumps(mapping, ensure_ascii=False, indent=2), encoding="utf-8")
    return structure_out, mapping_out


def load_word_mapping(mapping_path: Path | str) -> WordMapping:
    mapping_path = Path(mapping_path)
    if not mapping_path.exists():
        raise FileNotFoundError(f"Word mapping JSON not found: {mapping_path}")
    data = json.loads(mapping_path.read_text(encoding="utf-8"))

    def _get_int(*names: str, default: Optional[int] = None) -> int:
        for name in names:
            if name in data and data[name] is not None:
                return int(data[name])
        if default is not None:
            return default
        raise ValueError(f"Missing required mapping value: one of {', '.join(names)}")

    table_idx = _get_int("table_index_zero_based", "table_index", default=0)
    header_row = _get_int("header_row_zero_based", "header_row", default=0)
    data_start = _get_int("data_start_row_zero_based", "data_start_row", default=header_row + 1)
    cols = data.get("columns_by_index_zero_based") or data.get("columns_by_index") or {}
    clean_cols: Dict[str, int] = {}
    for field_name, idx in cols.items():
        if field_name not in WORD_CUSTOM_COMPARISON_FIELDS:
            continue
        if idx is None or str(idx).strip() == "":
            continue
        clean_cols[field_name] = int(idx)
    if not clean_cols:
        raise ValueError("Word mapping does not contain any usable columns. Edit columns_by_index_zero_based in the mapping JSON.")
    if "dok_id" not in clean_cols and not any(f in clean_cols for f in ("dokumentnummer", "artikelnummer")):
        raise ValueError("Word mapping should include DOK-ID, Dok.-Nr. or Artikel-Nr. so rows can be matched to Excel.")
    return WordMapping(
        table_index_zero_based=table_idx,
        header_row_zero_based=header_row,
        data_start_row_zero_based=data_start,
        columns_by_index_zero_based=clean_cols,
    )


def extract_word_records(docx_path: Path | str, mapping: WordMapping) -> List[WordRecord]:
    """Extract a custom Word table using an explicit mapping JSON."""
    docx_path = Path(docx_path)
    if not docx_path.exists():
        raise FileNotFoundError(f"Word file not found: {docx_path}")
    doc = Document(docx_path)
    if mapping.table_index_zero_based < 0 or mapping.table_index_zero_based >= len(doc.tables):
        raise ValueError(f"Word table index {mapping.table_index_zero_based} invalid. The document has {len(doc.tables)} table(s).")
    table = doc.tables[mapping.table_index_zero_based]
    records: List[WordRecord] = []
    for row_idx in range(mapping.data_start_row_zero_based, len(table.rows)):
        cells = table.rows[row_idx].cells
        raw: Dict[str, str] = {}
        norm: Dict[str, str] = {}
        for field_name, col_idx in mapping.columns_by_index_zero_based.items():
            value = _word_cell_to_text(cells[col_idx]) if 0 <= col_idx < len(cells) else ""
            raw[field_name] = value
            norm[field_name] = normalize_field(field_name, value)
        if any(raw.values()):
            records.append(WordRecord(
                word_file_name=docx_path.name,
                word_row_number=row_idx + 1,
                values_raw=raw,
                values_norm=norm,
            ))
    return records


def _cell_key(value: object) -> str:
    return normalize_header(value)


def _has_any(key: str, terms: Sequence[str]) -> bool:
    return any(term in key for term in terms)


def _unique_row_cells(row) -> List[Tuple[int, str, int]]:
    """Return rightmost grid index, text, and grid span for each distinct XML cell."""
    return [(last, text, span) for _first, last, _cell, text, span in _unique_row_cell_objects(row)]


def _unique_row_cell_objects(row) -> List[Tuple[int, int, Any, str, int]]:
    """Return first/last grid index, cell object, text and span for distinct cells."""
    seen: Dict[int, List[Any]] = {}
    for idx, cell in enumerate(row.cells):
        key = id(cell._tc)
        if key not in seen:
            seen[key] = [idx, idx, cell, _word_cell_to_text(cell)]
        else:
            seen[key][1] = idx
    result = [
        (first, last, cell, text, last - first + 1)
        for first, last, cell, text in seen.values()
    ]
    return sorted(result, key=lambda item: item[0])


def _clean_change_number_candidate(text: str) -> str:
    value = clean_text(text)
    value = re.sub(r"^\s*[\(\[]?\s*(?:number|nummer|no\.?|nr\.?)\s*[\)\]]?\s*[:\-]?\s*", "", value, flags=re.I)
    return value.strip()


def _paragraph_visible_text(paragraph) -> str:
    """Return all visible text from a Word paragraph, including content controls."""
    try:
        return clean_text("".join((node.text or "") for node in paragraph._p.xpath(".//w:t")))
    except Exception:
        return clean_text(getattr(paragraph, "text", ""))


def _extract_change_number_from_exact_header_address(doc) -> Tuple[
    str,
    List[str],
    Optional[int],
    Optional[int],
    Optional[int],
    Optional[int],
]:
    """Read the ECN number from its confirmed Word address.

    The user's template inspector identified the exact location as::

        section[0].header.table[0].row[0].cell[1].paragraph[0].word[0]

    ``python-docx`` has no direct ``word`` object, so the first whitespace-delimited
    token in paragraph 0 is used. Reading this exact header cell takes precedence
    over all heuristic/body-table detection.
    """
    warnings: List[str] = []
    section_idx = header_table_idx = row_idx = col_idx = 0
    col_idx = 1
    try:
        section = doc.sections[section_idx]
        header = section.header
        table = header.tables[header_table_idx]
        cell = table.rows[row_idx].cells[col_idx]
        if not cell.paragraphs:
            warnings.append(
                "The confirmed Word header Number cell exists but has no paragraph."
            )
            return "", warnings, section_idx, header_table_idx, row_idx, col_idx

        paragraph_text = _paragraph_visible_text(cell.paragraphs[0])
        words = paragraph_text.split()
        candidate = words[0] if words else ""
        cleaned = _clean_change_number_candidate(candidate)
        placeholder_key = _cell_key(cleaned)
        if not cleaned or placeholder_key in {"number", "nummer", "no", "nr"}:
            # A formatted/content-control implementation can occasionally put the
            # text outside the paragraph wrapper. Try the whole confirmed cell as
            # a narrow fallback, still without scanning any unrelated location.
            cell_text = _word_cell_to_text(cell)
            cell_words = cell_text.split()
            cleaned = _clean_change_number_candidate(cell_words[0] if cell_words else "")
            placeholder_key = _cell_key(cleaned)

        if not cleaned or placeholder_key in {"number", "nummer", "no", "nr"}:
            warnings.append(
                "The confirmed Word header Number address is empty or contains only its placeholder."
            )
            return "", warnings, section_idx, header_table_idx, row_idx, col_idx

        return cleaned, warnings, section_idx, header_table_idx, row_idx, col_idx
    except (IndexError, AttributeError) as exc:
        warnings.append(
            "The confirmed Word header Number address could not be opened: " + str(exc)
        )
        return "", warnings, section_idx, header_table_idx, row_idx, col_idx


def _set_first_word_in_confirmed_header_cell(cell, value: object) -> None:
    """Replace paragraph[0].word[0] while preserving the rest of the cell text."""
    desired = clean_text(value)
    if not cell.paragraphs:
        _set_word_cell_text_preserve_structure(cell, desired)
        return
    paragraph = cell.paragraphs[0]
    try:
        nodes = list(paragraph._p.xpath(".//w:t"))
    except Exception:
        nodes = []
    if not nodes:
        _set_word_cell_text_preserve_structure(cell, desired)
        return

    original = "".join((node.text or "") for node in nodes)
    if re.search(r"\S+", original):
        updated = re.sub(r"\S+", desired, original, count=1)
    else:
        updated = desired
    nodes[0].text = updated
    for node in nodes[1:]:
        node.text = ""


def _locate_change_notice_number_cell(
    table,
    header_start_row: int,
    title_rows: Optional[Sequence[int]] = None,
) -> Tuple[Optional[int], Optional[int], Optional[Any], List[str]]:
    """Locate the dedicated top-right Number cell in the standard template.

    The Number is not inferred from arbitrary codes elsewhere in the header. The
    standard template has a wide title cell and a separate narrow cell at its far
    right; that exact cell is used.
    """
    warnings: List[str] = []
    row_candidates: List[int] = []
    for idx in title_rows or []:
        if 0 <= idx < len(table.rows) and idx not in row_candidates:
            row_candidates.append(idx)
    for idx in range(max(0, min(header_start_row, len(table.rows)))):
        if idx not in row_candidates:
            row_candidates.append(idx)

    for row_idx in row_candidates:
        distinct = _unique_row_cell_objects(table.rows[row_idx])
        if len(distinct) < 2:
            continue
        # Prefer a row containing the title, then take its separate far-right cell.
        title_positions = [
            item for item in distinct
            if _has_any(_cell_key(item[3]), ["freigabeaenderungsmitteilung", "engineeringchangenotice"])
        ]
        if title_positions:
            title_last = max(item[1] for item in title_positions)
            right_cells = [item for item in distinct if item[0] > title_last]
            if right_cells:
                first, last, cell, _text, _span = max(right_cells, key=lambda item: item[1])
                return row_idx, last, cell, warnings

    # Structural fallback: first pre-header row with a broad left cell and a narrow
    # dedicated cell on the far right. This still targets the Number panel only.
    for row_idx in row_candidates:
        distinct = _unique_row_cell_objects(table.rows[row_idx])
        if len(distinct) < 2:
            continue
        left = min(distinct, key=lambda item: item[0])
        right = max(distinct, key=lambda item: item[1])
        if left[4] >= 2 and right[4] <= max(2, len(table.rows[row_idx].cells) // 4):
            return row_idx, right[1], right[2], warnings

    warnings.append("The dedicated top-right Number cell could not be located.")
    return None, None, None, warnings


def _extract_change_notice_number(
    table,
    header_start_row: int,
    title_rows: Optional[Sequence[int]] = None,
) -> Tuple[str, List[str], Optional[int], Optional[int]]:
    row_idx, col_idx, cell, warnings = _locate_change_notice_number_cell(
        table, header_start_row, title_rows=title_rows
    )
    if cell is None:
        return "", warnings, row_idx, col_idx
    raw = _word_cell_to_text(cell)
    cleaned = _clean_change_number_candidate(raw)
    placeholder_key = _cell_key(cleaned)
    if not cleaned or placeholder_key in {"number", "nummer", "no", "nr"}:
        warnings.append("The top-right Number cell is empty or still contains only its placeholder.")
        return "", warnings, row_idx, col_idx
    return cleaned, warnings, row_idx, col_idx


def _detect_change_notice_table(table, table_index: int) -> Optional[Dict[str, Any]]:
    """Recognize the bilingual Freigabe-/Änderungsmitteilung template by structure/header text."""
    rows = [[_word_cell_to_text(c) for c in row.cells] for row in table.rows]
    if not rows:
        return None
    max_cols = max(len(r) for r in rows)
    if max_cols < 9:
        return None

    article_hits: List[Tuple[int, int]] = []
    doc_hits: List[Tuple[int, int]] = []
    rev_new_hits: List[Tuple[int, int]] = []
    vers_new_hits: List[Tuple[int, int]] = []
    sheets_hits: List[Tuple[int, int]] = []
    title_hits: List[Tuple[int, int]] = []

    for r_idx, row in enumerate(rows[: min(12, len(rows))]):
        for c_idx, text in enumerate(row):
            key = _cell_key(text)
            if not key:
                continue
            if _has_any(key, ["freigabeaenderungsmitteilung", "engineeringchangenotice"]):
                title_hits.append((r_idx, c_idx))
            if ("artikel" in key and "nr" in key) or "partno" in key or "partnumber" in key:
                article_hits.append((r_idx, c_idx))
            if _has_any(key, ["dokumentnr", "dokumentnummer", "documentno", "documentnumber", "doknr"]):
                doc_hits.append((r_idx, c_idx))
            if _has_any(key, ["revneunew", "revnew", "revneu"]):
                rev_new_hits.append((r_idx, c_idx))
            if _has_any(key, ["versneunew", "versnew", "versneu", "versionnew"]):
                vers_new_hits.append((r_idx, c_idx))
            if _has_any(key, ["blattsheets", "sheets", "blatt", "pages", "seiten"]):
                sheets_hits.append((r_idx, c_idx))

    if not article_hits or not doc_hits or not rev_new_hits or not vers_new_hits:
        return None

    doc_row, doc_col = min(doc_hits, key=lambda item: (item[0], item[1]))
    article_before = [hit for hit in article_hits if hit[1] < doc_col]
    article_row, article_col = min(article_before or article_hits, key=lambda item: (item[0], item[1]))
    rev_after = [hit for hit in rev_new_hits if hit[1] > doc_col]
    vers_after = [hit for hit in vers_new_hits if hit[1] > doc_col]
    if not rev_after or not vers_after:
        return None
    rev_row, rev_col = min(rev_after, key=lambda item: item[1])
    vers_row, vers_col = min(vers_after, key=lambda item: item[1])
    if rev_col == vers_col:
        return None
    sheets_after = [hit for hit in sheets_hits if hit[1] > vers_col]
    sheets_row, sheets_col = min(sheets_after, key=lambda item: item[1]) if sheets_after else (-1, -1)

    extra_rows = [sheets_row] if sheets_row >= 0 else []
    header_rows = sorted(set([article_row, doc_row, rev_row, vers_row] + extra_rows + [r for r, _ in title_hits]))
    detail_header_rows = sorted(set([article_row, doc_row, rev_row, vers_row] + extra_rows))
    data_start = max(detail_header_rows) + 1
    while data_start < len(rows):
        row = rows[data_start]
        check_cols = [article_col, doc_col, rev_col, vers_col] + ([sheets_col] if sheets_col >= 0 else [])
        if any(clean_text(row[c]) for c in check_cols if c < len(row)):
            break
        data_start += 1

    change_raw, warnings, change_row, change_col = _extract_change_notice_number(
        table, min(detail_header_rows), title_rows=[r for r, _ in title_hits]
    )
    score = 100 + (20 if title_hits else 0) + min(max_cols, 20)
    return {
        "score": score,
        "table_index": table_index,
        "header_rows": header_rows,
        "data_start": data_start,
        "columns": {
            "artikelnummer": article_col,
            "dokumentnummer": doc_col,
            "revision": rev_col,
            "version": vers_col,
            **({"seiten": sheets_col} if sheets_col >= 0 else {}),
        },
        "change_number_raw": change_raw,
        "change_number_cell_row": change_row,
        "change_number_cell_col": change_col,
        "warnings": warnings + ([] if sheets_col >= 0 else ["The Blatt / sheets column was not detected; PDF page-count comparison will be unavailable."]),
    }


def extract_change_notice_word_records(docx_path: Path | str) -> Tuple[List[WordRecord], WordExtractionInfo]:
    docx_path = Path(docx_path)
    if not docx_path.exists():
        raise FileNotFoundError(f"Word file not found: {docx_path}")
    doc = Document(docx_path)
    detections: List[Tuple[Dict[str, Any], Any]] = []
    for table_idx, table in enumerate(doc.tables):
        detected = _detect_change_notice_table(table, table_idx)
        if detected:
            detections.append((detected, table))
    if not detections:
        raise ValueError(
            "Word table not recognized as the Freigabe-/Änderungsmitteilung / Engineering Change Notice template. "
            "Choose a custom mapping JSON only if this is a different table template."
        )
    detections.sort(key=lambda item: item[0]["score"], reverse=True)
    detected, table = detections[0]
    columns = detected["columns"]
    (
        exact_change_raw,
        exact_warnings,
        exact_section_idx,
        exact_header_table_idx,
        exact_row_idx,
        exact_col_idx,
    ) = _extract_change_number_from_exact_header_address(doc)

    # The confirmed section/header address is authoritative. The prior body-table
    # locator remains only as a backward-compatible fallback for older synthetic
    # documents and variants that do not contain the expected header structure.
    if exact_change_raw:
        change_raw = exact_change_raw
        change_source = "section[0].header.table[0].row[0].cell[1].paragraph[0].word[0]"
        change_section_idx = exact_section_idx
        change_header_table_idx = exact_header_table_idx
        change_row_idx = exact_row_idx
        change_col_idx = exact_col_idx
        change_warnings = exact_warnings
    else:
        change_raw = detected["change_number_raw"]
        change_source = "body_table_fallback"
        change_section_idx = None
        change_header_table_idx = None
        change_row_idx = detected.get("change_number_cell_row")
        change_col_idx = detected.get("change_number_cell_col")
        change_warnings = exact_warnings + [
            "The confirmed header address yielded no usable Number; the legacy body-table locator was used."
        ]
    change_norm = normalize_field("freigabe", change_raw)
    records: List[WordRecord] = []
    for row_idx in range(detected["data_start"], len(table.rows)):
        cells = table.rows[row_idx].cells
        raw: Dict[str, str] = {}
        for field_name, col_idx in columns.items():
            raw[field_name] = _word_cell_to_text(cells[col_idx]) if 0 <= col_idx < len(cells) else ""
        # The top-right Number applies globally to every listed document.
        raw["freigabe"] = change_raw
        norm = {field_name: normalize_field(field_name, value) for field_name, value in raw.items()}
        identity_values = [raw.get("dokumentnummer", ""), raw.get("artikelnummer", ""), raw.get("revision", ""), raw.get("version", "")]
        if not any(identity_values):
            continue
        # Do not accidentally treat repeated headers as data.
        if _has_any(_cell_key(raw.get("dokumentnummer", "")), ["dokumentnr", "documentno"]):
            continue
        records.append(WordRecord(
            word_file_name=docx_path.name,
            word_row_number=row_idx + 1,
            values_raw=raw,
            values_norm=norm,
        ))
    info = WordExtractionInfo(
        mode="automatic_template",
        template_type="Freigabe-/Änderungsmitteilung / Engineering Change Notice",
        table_index_zero_based=detected["table_index"],
        header_rows_zero_based=detected["header_rows"],
        data_start_row_zero_based=detected["data_start"],
        columns_by_index_zero_based=dict(columns),
        change_number_raw=change_raw,
        change_number_norm=change_norm,
        change_number_source=change_source,
        change_number_section_zero_based=change_section_idx,
        change_number_header_table_zero_based=change_header_table_idx,
        change_number_cell_row_zero_based=change_row_idx,
        change_number_cell_col_zero_based=change_col_idx,
        warnings=list(detected["warnings"]) + change_warnings,
    )
    return records, info


def preview_word_extraction(docx_path: Path | str, mapping_path: Optional[Path | str] = None, limit: int = 10) -> Tuple[List[WordRecord], WordExtractionInfo]:
    if mapping_path:
        mapping = load_word_mapping(mapping_path)
        records = extract_word_records(docx_path, mapping)
        info = WordExtractionInfo(
            mode="custom_mapping",
            template_type="Custom Word mapping",
            table_index_zero_based=mapping.table_index_zero_based,
            header_rows_zero_based=[mapping.header_row_zero_based],
            data_start_row_zero_based=mapping.data_start_row_zero_based,
            columns_by_index_zero_based=dict(mapping.columns_by_index_zero_based),
        )
        return records[:limit], info
    records, info = extract_change_notice_word_records(docx_path)
    return records[:limit], info


def _score_word_candidate(word_record: WordRecord, record: ExcelRecord) -> float:
    earned = 0
    possible = 0
    for field_name, weight in WORD_MATCH_WEIGHTS.items():
        if field_name not in word_record.values_norm:
            continue
        word_norm = word_record.values_norm.get(field_name, "")
        xl_norm = record.values_norm.get(field_name, "")
        if not word_norm or not xl_norm:
            continue
        possible += weight
        if word_norm == xl_norm:
            earned += weight
    if possible == 0:
        return 0.0
    return round(100.0 * earned / possible, 1)


def _word_candidate_list(word_record: WordRecord, records: Sequence[ExcelRecord]) -> List[Tuple[ExcelRecord, float]]:
    scored = [(rec, _score_word_candidate(word_record, rec)) for rec in records]
    scored.sort(key=lambda x: x[1], reverse=True)
    return scored


def compare_word_value(field_name: str, word_raw: str, excel_raw: str, word_file_name: str, word_row_number: int, excel_row: Optional[int]) -> WordFieldComparison:
    word_clean = clean_text(word_raw)
    excel_clean = clean_text(excel_raw)
    word_norm = normalize_field(field_name, word_raw)
    excel_norm = normalize_field(field_name, excel_raw)
    if not word_clean and not excel_clean:
        res = "BOTH_EMPTY"
    elif not word_clean:
        res = "WORD_MISSING"
    elif not excel_clean:
        res = "EXCEL_MISSING"
    elif field_name == "sprache":
        res = "OK" if word_norm and word_norm == excel_norm else "MISMATCH"
    else:
        res = "OK" if word_clean == excel_clean else "MISMATCH"
    return WordFieldComparison(
        word_file_name=word_file_name,
        word_row_number=word_row_number,
        field=PDF_FIELD_DISPLAY[field_name],
        word_raw=word_clean,
        word_norm=word_norm,
        excel_raw=excel_clean,
        excel_norm=excel_norm,
        result=res,
        excel_row=excel_row,
    )


def validate_one_word_record(
    word_record: WordRecord,
    records: Sequence[ExcelRecord],
    dok_id_index: Dict[str, List[ExcelRecord]],
) -> WordValidationResult:
    result = WordValidationResult(record=word_record, status="WORD_NO_RELIABLE_MATCH")
    matched_record: Optional[ExcelRecord] = None
    match_status: Optional[str] = None

    # Custom mappings may still provide DOK-ID; use it when available.
    word_dok = word_record.values_norm.get("dok_id", "")
    wrong_dok_id = bool(word_dok and word_dok not in dok_id_index)
    if word_dok and word_dok in dok_id_index:
        rows = dok_id_index[word_dok]
        if len(rows) == 1:
            matched_record = rows[0]
            result.confidence = 100.0
        else:
            scored = _word_candidate_list(word_record, rows)
            result.candidate_rows = [(rec.row_number, score) for rec, score in scored[:10]]
            if _best_candidate_is_clear(scored, threshold=70.0, margin=10.0):
                matched_record = scored[0][0]
                result.confidence = scored[0][1]
                match_status = "WORD_DUPLICATE_DOCUMENT_NUMBER_RESOLVED"
                result.issues.append(f"DOK-ID appears {len(rows)} times in Excel; the best row was selected using the other values.")
            else:
                result.status = "WORD_AMBIGUOUS_MATCH"
                result.confidence = scored[0][1] if scored else 0.0
                result.issues.append(f"DOK-ID appears {len(rows)} times in Excel and no unique best row was found.")
                return result

    if matched_record is None:
        word_doc = word_record.values_norm.get("dokumentnummer", "")
        exact_doc_rows = [rec for rec in records if word_doc and rec.values_norm.get("dokumentnummer", "") == word_doc]
        if len(exact_doc_rows) == 1:
            matched_record = exact_doc_rows[0]
            result.confidence = 100.0
            if wrong_dok_id:
                match_status = "WORD_LIKELY_WRONG_DOK_ID"
                result.issues.append("The Word DOK-ID is not present in Excel, but Dokument-Nr. uniquely identifies the row.")
        elif len(exact_doc_rows) > 1:
            scored = _word_candidate_list(word_record, exact_doc_rows)
            result.candidate_rows = [(rec.row_number, score) for rec, score in scored[:10]]
            if _best_candidate_is_clear(scored, threshold=70.0, margin=8.0):
                matched_record = scored[0][0]
                result.confidence = scored[0][1]
                match_status = "WORD_DUPLICATE_DOCUMENT_NUMBER_RESOLVED"
                result.issues.append(f"Dokument-Nr. appears {len(exact_doc_rows)} times in Excel; Artikel-Nr./revision/version resolved it.")
            else:
                result.status = "WORD_AMBIGUOUS_MATCH"
                result.confidence = scored[0][1] if scored else 0.0
                result.issues.append(f"Dokument-Nr. appears {len(exact_doc_rows)} times in Excel and no unique best row was found.")
                return result
        else:
            scored = _word_candidate_list(word_record, records)
            result.candidate_rows = [(rec.row_number, score) for rec, score in scored[:10]]
            if _best_candidate_is_clear(scored, threshold=70.0, margin=15.0):
                matched_record = scored[0][0]
                result.confidence = scored[0][1]
                match_status = "WORD_LIKELY_WRONG_DOCUMENT_NUMBER"
                result.issues.append("The Word Dokument-Nr. is missing or not in Excel, but Artikel-Nr./revision/version strongly identify another Excel row.")
            else:
                result.confidence = scored[0][1] if scored else 0.0
                result.status = "WORD_DOCUMENT_NOT_IN_EXCEL" if word_doc else "WORD_NO_RELIABLE_MATCH"
                result.issues.append("The Word row could not be reliably matched to Excel.")
                return result

    result.matched_excel_row = matched_record.row_number if matched_record else None
    comps: List[WordFieldComparison] = []
    if matched_record:
        for field_name in WORD_EXCEL_COMPARISON_FIELDS:
            if field_name not in word_record.values_raw:
                continue
            comps.append(compare_word_value(
                field_name,
                word_record.values_raw.get(field_name, ""),
                matched_record.values_raw.get(field_name, ""),
                word_file_name=word_record.word_file_name,
                word_row_number=word_record.word_row_number,
                excel_row=matched_record.row_number,
            ))
    result.comparisons = comps
    bad_change = [c for c in comps if c.field == PDF_FIELD_DISPLAY["freigabe"] and c.result not in {"OK", "BOTH_EMPTY"}]
    bad_other = [c for c in comps if c.field != PDF_FIELD_DISPLAY["freigabe"] and c.result not in {"OK", "BOTH_EMPTY"}]

    if match_status == "WORD_LIKELY_WRONG_DOK_ID" and not bad_change and bad_other and all(c.field == PDF_FIELD_DISPLAY["dok_id"] for c in bad_other):
        result.status = match_status
    elif bad_change and bad_other:
        result.status = "WORD_FIELD_AND_CHANGE_NUMBER_MISMATCH"
        result.issues.append("Word/Excel field mismatches: " + ", ".join(c.field for c in bad_other))
        result.issues.append("The top-right Word Number differs from Excel Freigabe-/Änd.-Nr.")
    elif bad_change:
        result.status = "WORD_CHANGE_NUMBER_MISMATCH"
        result.issues.append("The top-right Word Number differs from Excel Freigabe-/Änd.-Nr.")
    elif bad_other:
        result.status = "WORD_FIELD_MISMATCH"
        result.issues.append("Word/Excel field mismatches: " + ", ".join(c.field for c in bad_other))
    elif match_status:
        result.status = match_status
    else:
        result.status = "WORD_OK"
    return result


def validate_word_records(
    word_records: Sequence[WordRecord],
    records: Sequence[ExcelRecord],
    dok_id_index: Dict[str, List[ExcelRecord]],
) -> List[WordValidationResult]:
    return [validate_one_word_record(rec, records, dok_id_index) for rec in word_records]


def apply_word_page_count_checks(
    word_results: Sequence[WordValidationResult],
    pdf_results: Sequence[ValidationResult],
    pdfs_available: bool,
) -> None:
    """Compare Word Blatt/sheets values with actual PDF page counts.

    Word rows and PDFs are linked through the Excel row selected by the existing
    matching logic. This avoids relying on filenames.
    """
    by_excel_row: Dict[int, List[ValidationResult]] = {}
    for pdf_result in pdf_results:
        if pdf_result.matched_excel_row and pdf_result.extracted.total_pages is not None:
            by_excel_row.setdefault(pdf_result.matched_excel_row, []).append(pdf_result)

    non_error_excel_statuses = {
        "WORD_OK", "WORD_OK_WITH_NORMALIZATION",
        "WORD_DUPLICATE_DOCUMENT_NUMBER_RESOLVED",
        "WORD_LIKELY_WRONG_DOCUMENT_NUMBER", "WORD_LIKELY_WRONG_DOK_ID",
    }

    for result in word_results:
        result.excel_status = result.status
        result.word_pages_raw = clean_text(result.record.values_raw.get("seiten", ""))
        result.word_pages_norm = normalize_page_count(result.word_pages_raw)

        if not pdfs_available:
            result.page_count_result = "UNAVAILABLE_NO_PDFS"
            result.page_count_issue = "No PDF folder was provided, so the true number of pages could not be verified."
            result.issues.append(result.page_count_issue)
            continue
        if not result.matched_excel_row:
            result.page_count_result = "UNAVAILABLE_NO_EXCEL_MATCH"
            result.page_count_issue = "The Word row was not matched to Excel, so it could not be linked to a PDF."
            result.issues.append(result.page_count_issue)
            continue
        if not result.word_pages_raw:
            result.page_count_result = "WORD_PAGE_COUNT_MISSING"
            result.page_count_issue = "The Word Blatt / sheets cell is empty."
            result.issues.append(result.page_count_issue)
            continue
        if not result.word_pages_norm:
            result.page_count_result = "WORD_PAGE_COUNT_UNSUPPORTED"
            result.page_count_issue = f"The Word Blatt / sheets value could not be interpreted: {result.word_pages_raw!r}."
            result.issues.append(result.page_count_issue)
            continue

        candidates = by_excel_row.get(result.matched_excel_row, [])
        if not candidates:
            result.page_count_result = "PDF_NOT_FOUND"
            result.page_count_issue = "No successfully matched PDF was found for this Excel row."
            result.issues.append(result.page_count_issue)
            continue

        counts = sorted({c.extracted.total_pages for c in candidates if c.extracted.total_pages is not None})
        files = sorted(c.extracted.file_name for c in candidates)
        result.pdf_file_name = "; ".join(files)
        if len(counts) > 1:
            result.page_count_result = "AMBIGUOUS_PDF_PAGE_COUNT"
            result.page_count_issue = (
                "Several PDFs matched the same Excel row but have different page counts: "
                + ", ".join(str(v) for v in counts)
            )
            result.issues.append(result.page_count_issue)
            continue

        result.pdf_pages = counts[0]
        if int(result.word_pages_norm) == result.pdf_pages:
            result.page_count_result = "OK"
            if len(candidates) > 1:
                result.page_count_issue = "Several PDFs matched this Excel row, but all have the same page count."
                result.issues.append(result.page_count_issue)
        else:
            result.page_count_result = "MISMATCH"
            result.page_count_issue = (
                f"Word Blatt / sheets is {result.word_pages_raw}, but the matched PDF has {result.pdf_pages} page(s)."
            )
            result.issues.append(result.page_count_issue)
            if result.excel_status in non_error_excel_statuses:
                result.status = "WORD_PAGE_COUNT_MISMATCH"
            else:
                result.status = "WORD_MULTIPLE_MISMATCHES"


def run_validation(
    pdf_folder: Optional[Path | str],
    excel_path: Path | str,
    output_path: Path | str,
    expected_ausgabe: Optional[str],
    sheet_name: Optional[str] = None,
    page_number: int = 2,
    table_index: Optional[int] = None,
    row_start: Optional[int] = None,
    word_docx_path: Optional[Path | str] = None,
    word_mapping_path: Optional[Path | str] = None,
    progress_callback=None,
    summary_callback=None,
) -> List[ValidationResult]:
    has_pdf_folder = bool(clean_text(pdf_folder))
    has_word = bool(clean_text(word_docx_path))
    if not has_pdf_folder and not has_word:
        raise ValueError("Choose a PDF folder, a Word file, or both.")

    expected_input = clean_text(expected_ausgabe)
    expected_norm = ""
    if has_pdf_folder:
        if not expected_input:
            raise ValueError("Expected Ausgabe is required when a PDF folder is selected.")
        expected_norm = normalize_expected_ausgabe(expected_input)

    if progress_callback:
        progress_callback("Loading Excel...")
    records, actual_sheet, header_row, columns = load_excel_records(excel_path, sheet_name)
    if not records:
        raise ValueError("No data rows found in Excel after the required header row.")
    if progress_callback:
        progress_callback(f"Loaded {len(records)} Excel rows from sheet '{actual_sheet}' using header row {header_row}.")
    dok_idx = build_dok_id_index(records)

    results: List[ValidationResult] = []
    pdfs: List[Path] = []
    if has_pdf_folder:
        pdfs = list_pdf_files(pdf_folder)
        if not pdfs:
            raise ValueError("No PDF files found in the selected folder.")
        if progress_callback:
            progress_callback(f"Found {len(pdfs)} PDF files.")
        for i, pdf in enumerate(pdfs, start=1):
            if progress_callback:
                progress_callback(f"Processing PDF {i}/{len(pdfs)}: {pdf.name}")
            extracted = extract_pdf_table(pdf, page_number=page_number, table_index=table_index, row_start=row_start)
            res = validate_one_pdf(extracted, records, dok_idx, expected_input, expected_norm)
            results.append(res)
    elif progress_callback:
        progress_callback("No PDF folder selected. PDF metadata and true page counts will not be checked.")

    word_records: Optional[List[WordRecord]] = None
    word_results: Optional[List[WordValidationResult]] = None
    word_mapping: Optional[WordMapping] = None
    word_info: Optional[WordExtractionInfo] = None
    if has_word:
        if progress_callback:
            if word_mapping_path:
                progress_callback("Loading Word table with custom mapping...")
            else:
                progress_callback("Recognizing the Freigabe-/Änderungsmitteilung Word template...")
        if word_mapping_path:
            word_mapping = load_word_mapping(word_mapping_path)
            word_records = extract_word_records(word_docx_path, word_mapping)
            word_info = WordExtractionInfo(
                mode="custom_mapping",
                template_type="Custom Word mapping",
                table_index_zero_based=word_mapping.table_index_zero_based,
                header_rows_zero_based=[word_mapping.header_row_zero_based],
                data_start_row_zero_based=word_mapping.data_start_row_zero_based,
                columns_by_index_zero_based=dict(word_mapping.columns_by_index_zero_based),
            )
        else:
            word_records, word_info = extract_change_notice_word_records(word_docx_path)
        if not word_records:
            raise ValueError("No document rows were found in the selected Word table.")
        if progress_callback:
            progress_callback(f"Loaded {len(word_records)} Word rows from '{Path(word_docx_path).name}'.")
            if word_info and word_info.change_number_raw:
                progress_callback(f"Word change-notice Number: {word_info.change_number_raw}")
            progress_callback("Comparing Word rows with Excel...")
        word_results = validate_word_records(word_records, records, dok_idx)
        apply_word_page_count_checks(word_results, results, pdfs_available=has_pdf_folder)
        if progress_callback:
            word_counts: Dict[str, int] = {}
            for word_result in word_results:
                word_counts[word_result.status] = word_counts.get(word_result.status, 0) + 1
            progress_callback(
                "Word results: "
                + ", ".join(
                    f"{status}={count}"
                    for status, count in sorted(
                        word_counts.items(), key=lambda item: WORD_STATUS_ORDER.get(item[0], 99)
                    )
                )
            )
    elif word_mapping_path:
        raise ValueError("A Word mapping JSON was selected without a Word .docx file.")

    if progress_callback:
        progress_callback("Writing report...")
    write_report(
        results,
        records,
        output_path,
        expected_input,
        expected_norm,
        actual_sheet,
        word_results=word_results,
        word_records=word_records,
        word_mapping=word_mapping,
        word_extraction_info=word_info,
        word_docx_path=word_docx_path,
        pdf_folder_provided=has_pdf_folder,
        excel_path=excel_path,
        pdf_folder=pdf_folder,
    )
    if summary_callback:
        summary_callback(build_compact_run_summary(
            results, records, word_results,
            pdf_folder_provided=has_pdf_folder,
            word_file_provided=has_word,
        ))
    if progress_callback:
        progress_callback(f"Done. Report saved to: {output_path}")
    return results


def preview_extraction(
    pdf_folder: Path | str,
    page_number: int = 2,
    table_index: Optional[int] = None,
    row_start: Optional[int] = None,
    limit: int = 5,
) -> List[ExtractedPDF]:
    pdfs = list_pdf_files(pdf_folder)[:limit]
    return [extract_pdf_table(pdf, page_number=page_number, table_index=table_index, row_start=row_start) for pdf in pdfs]


def _join_unique(values: Iterable[object], separator: str = "; ") -> str:
    seen: List[str] = []
    for value in values:
        text = clean_text(value)
        if text and text not in seen:
            seen.append(text)
    return separator.join(seen)


PDF_STATUS_LABELS = {
    "OK": "OK",
    "OK_WITH_NORMALIZATION": "OK",
    "AUSGABE_MISMATCH": "Ausgabe mismatch",
    "MISMATCH": "Field mismatch",
    "MISMATCH_AND_AUSGABE_MISMATCH": "Field and Ausgabe mismatch",
    "LIKELY_WRONG_DOK_ID": "Likely wrong DOK-ID",
    "DUPLICATE_DOK_ID_RESOLVED": "Duplicate DOK-ID resolved",
    "DUPLICATE_DOK_ID_AMBIGUOUS": "Duplicate DOK-ID ambiguous",
    "DOK_ID_NOT_FOUND": "DOK-ID not found",
    "NO_RELIABLE_MATCH": "No reliable Excel match",
    "EXTRACTION_ERROR": "PDF extraction error",
}

WORD_STATUS_LABELS = {
    "WORD_OK": "OK",
    "WORD_OK_WITH_NORMALIZATION": "OK",
    "WORD_DUPLICATE_DOCUMENT_NUMBER_RESOLVED": "Duplicate document number resolved",
    "WORD_LIKELY_WRONG_DOCUMENT_NUMBER": "Likely wrong document number",
    "WORD_LIKELY_WRONG_DOK_ID": "Likely wrong DOK-ID",
    "WORD_CHANGE_NUMBER_MISMATCH": "Change number mismatch",
    "WORD_FIELD_MISMATCH": "Field mismatch",
    "WORD_FIELD_AND_CHANGE_NUMBER_MISMATCH": "Field and change-number mismatch",
    "WORD_DOCUMENT_NOT_IN_EXCEL": "Document not found in Excel",
    "WORD_AMBIGUOUS_MATCH": "Ambiguous Excel match",
    "WORD_NO_RELIABLE_MATCH": "No reliable Excel match",
    "WORD_TABLE_NOT_RECOGNIZED": "Word table not recognized",
    "WORD_PAGE_COUNT_MISMATCH": "Page-count mismatch",
    "WORD_MULTIPLE_MISMATCHES": "Multiple mismatches",
    "WORD_EXTRACTION_ERROR": "Word extraction error",
}

PAGE_RESULT_LABELS = {
    "OK": "OK",
    "NOT_CHECKED": "Not checked",
    "UNAVAILABLE_NO_PDFS": "Unavailable — no PDFs",
    "UNAVAILABLE_NO_EXCEL_MATCH": "Unavailable — no Excel match",
    "WORD_PAGE_COUNT_MISSING": "Word page count missing",
    "WORD_PAGE_COUNT_UNSUPPORTED": "Unsupported Word page count",
    "PDF_NOT_FOUND": "PDF not found",
    "AMBIGUOUS_PDF_PAGE_COUNT": "Ambiguous PDF page count",
    "MISMATCH": "Mismatch",
}


def _severity_rank(value: str) -> int:
    return {"OK": 0, "WARNING": 1, "ERROR": 2}.get(value, 2)


def _max_severity(values: Iterable[str]) -> str:
    best = "OK"
    for value in values:
        if _severity_rank(value) > _severity_rank(best):
            best = value
    return best


def _pdf_severity(status: str) -> str:
    if status in {"OK", "OK_WITH_NORMALIZATION"}:
        return "OK"
    if status in {"LIKELY_WRONG_DOK_ID", "DUPLICATE_DOK_ID_RESOLVED"}:
        return "WARNING"
    return "ERROR"


def _word_severity(result: WordValidationResult) -> str:
    if result.status in {"WORD_OK", "WORD_OK_WITH_NORMALIZATION"}:
        base = "OK"
    elif result.status in {
        "WORD_DUPLICATE_DOCUMENT_NUMBER_RESOLVED",
        "WORD_LIKELY_WRONG_DOCUMENT_NUMBER",
        "WORD_LIKELY_WRONG_DOK_ID",
    }:
        base = "WARNING"
    else:
        base = "ERROR"
    if result.page_count_result == "UNAVAILABLE_NO_PDFS":
        return _max_severity([base, "WARNING"])
    if result.page_count_result not in {"OK", "NOT_CHECKED", ""}:
        return _max_severity([base, "ERROR"])
    return base


def build_compact_run_summary(
    pdf_results: Sequence[ValidationResult],
    excel_records: Sequence[ExcelRecord],
    word_results: Optional[Sequence[WordValidationResult]],
    pdf_folder_provided: bool,
    word_file_provided: bool,
) -> Dict[str, Any]:
    """Build the concise result model shown in the application's right panel.

    Only actionable exceptions are listed. Each item contains the reliable Excel
    DOK-ID when one is available and a short human-readable description of the
    incorrect value.
    """
    excel_by_row = {record.row_number: record for record in excel_records}
    issues: List[Dict[str, str]] = []
    seen: set[Tuple[str, str, str]] = set()

    def add(severity: str, dok_id: str, issue: str, source: str = "") -> None:
        key = (severity, clean_text(dok_id), clean_text(issue))
        if not issue or key in seen:
            return
        seen.add(key)
        issues.append({
            "severity": severity,
            "dok_id": clean_text(dok_id),
            "issue": clean_text(issue),
            "source": clean_text(source),
        })

    pdf_by_excel: Dict[int, List[ValidationResult]] = {}
    for result in pdf_results:
        if result.matched_excel_row:
            pdf_by_excel.setdefault(result.matched_excel_row, []).append(result)
        record = excel_by_row.get(result.matched_excel_row or -1)
        reliable_id = record.values_raw.get("dok_id", "") if record else ""
        dok_comparison = next((c for c in result.comparisons if c.field == PDF_FIELD_DISPLAY["dok_id"]), None)
        if dok_comparison and dok_comparison.result not in {"OK", "BOTH_EMPTY"}:
            reliable_id = ""  # The displayed ID itself is the incorrect value.
        for comp in result.comparisons:
            if comp.result in {"OK", "BOTH_EMPTY"}:
                continue
            severity = "WARNING" if result.status in {"LIKELY_WRONG_DOK_ID", "DUPLICATE_DOK_ID_RESOLVED"} and comp.field == PDF_FIELD_DISPLAY["dok_id"] else "ERROR"
            add(
                severity,
                reliable_id,
                f'{comp.field}: PDF "{comp.pdf_raw or "<empty>"}" ≠ {comp.source} "{comp.reference_raw or "<empty>"}"',
                result.extracted.file_name,
            )
        if result.status not in {"OK", "OK_WITH_NORMALIZATION", "MISMATCH", "AUSGABE_MISMATCH", "MISMATCH_AND_AUSGABE_MISMATCH"}:
            add(_pdf_severity(result.status), reliable_id, _join_unique(result.issues) or PDF_STATUS_LABELS.get(result.status, result.status), result.extracted.file_name)
        for warning in result.extracted.warnings:
            add("WARNING", reliable_id, warning, result.extracted.file_name)

    word_by_excel: Dict[int, List[WordValidationResult]] = {}
    for result in word_results or []:
        if result.matched_excel_row:
            word_by_excel.setdefault(result.matched_excel_row, []).append(result)
        record = excel_by_row.get(result.matched_excel_row or -1)
        reliable_id = record.values_raw.get("dok_id", "") if record else ""
        for comp in result.comparisons:
            if comp.result in {"OK", "BOTH_EMPTY"}:
                continue
            add(
                _word_severity(result),
                reliable_id,
                f'{comp.field}: Word "{comp.word_raw or "<empty>"}" ≠ Excel "{comp.excel_raw or "<empty>"}"',
                f"Word row {result.record.word_row_number}",
            )
        if result.page_count_result not in {"OK", "NOT_CHECKED", ""}:
            if result.page_count_result == "UNAVAILABLE_NO_PDFS":
                add("WARNING", reliable_id, "Blatt / sheets could not be verified because no PDF folder was selected.", f"Word row {result.record.word_row_number}")
            else:
                pdf_value = str(result.pdf_pages) if result.pdf_pages is not None else "unavailable"
                add(
                    "ERROR",
                    reliable_id,
                    f'Blatt / sheets: Word "{result.word_pages_raw or "<empty>"}" ≠ PDF "{pdf_value}"',
                    f"Word row {result.record.word_row_number}",
                )
        if result.status in {
            "WORD_DOCUMENT_NOT_IN_EXCEL", "WORD_AMBIGUOUS_MATCH", "WORD_NO_RELIABLE_MATCH",
            "WORD_TABLE_NOT_RECOGNIZED", "WORD_EXTRACTION_ERROR",
        }:
            add(_word_severity(result), reliable_id, _join_unique(result.issues) or WORD_STATUS_LABELS.get(result.status, result.status), f"Word row {result.record.word_row_number}")

    # Show expected source presence for each Excel entry, just like the workbook report.
    for record in excel_records:
        dok_id = record.values_raw.get("dok_id", "")
        document = record.values_raw.get("dokumentnummer", "")
        if pdf_folder_provided and record.row_number not in pdf_by_excel:
            add("ERROR", dok_id, f'PDF not found for Dokumentnummer "{document}".', "PDF folder")
        if word_file_provided and record.row_number not in word_by_excel:
            add("ERROR", dok_id, f'Word entry not found for Dokumentnummer "{document}".', "Word table")

    issues.sort(key=lambda item: (-_severity_rank(item["severity"]), item["dok_id"], item["issue"]))
    error_count = sum(1 for item in issues if item["severity"] == "ERROR")
    warning_count = sum(1 for item in issues if item["severity"] == "WARNING")
    all_correct = not issues
    word_correction_needed = False
    for result in word_results or []:
        if any(comp.result not in {"OK", "BOTH_EMPTY"} for comp in result.comparisons):
            word_correction_needed = True
            break
        if result.page_count_result in {
            "WORD_PAGE_COUNT_MISSING", "WORD_PAGE_COUNT_UNSUPPORTED",
            "PDF_NOT_FOUND", "AMBIGUOUS_PDF_PAGE_COUNT", "MISMATCH",
        }:
            word_correction_needed = True
            break
    return {
        "all_correct": all_correct,
        "checked_entries": len(excel_records),
        "issue_count": len(issues),
        "error_count": error_count,
        "warning_count": warning_count,
        "issues": issues,
        "word_correction_available": bool(word_file_provided),
        "word_correction_needed": word_correction_needed,
    }


def create_corrected_word_copy(
    word_docx_path: Path | str,
    excel_path: Path | str,
    output_path: Path | str,
    sheet_name: Optional[str] = None,
    pdf_folder: Optional[Path | str] = None,
    progress_callback=None,
) -> WordCorrectionResult:
    """Create a corrected copy of the standard Word change-notice document.

    The original file is never overwritten. Matched Word rows receive exact Excel
    values for Artikel-Nr., Dokument-Nr., Rev neu and Vers neu. The dedicated
    top-right Number receives the common Excel Freigabe-/Änd.-Nr. When PDFs are
    available, Blatt / sheets receives the actual PDF page count.
    """
    word_docx_path = Path(word_docx_path)
    output_path = Path(output_path)
    if word_docx_path.resolve() == output_path.resolve():
        raise ValueError("Choose a different output path; the original Word file is not overwritten.")
    if progress_callback:
        progress_callback("Loading Excel and matching Word rows...")
    excel_records, _sheet, _header, _columns = load_excel_records(excel_path, sheet_name)
    excel_by_row = {record.row_number: record for record in excel_records}
    dok_idx = build_dok_id_index(excel_records)
    word_records, info = extract_change_notice_word_records(word_docx_path)
    word_results = validate_word_records(word_records, excel_records, dok_idx)

    pdf_results: List[ValidationResult] = []
    if clean_text(pdf_folder):
        pdf_files = list_pdf_files(pdf_folder)
        for pdf in pdf_files:
            extracted = extract_pdf_table(pdf, page_number=2)
            # Ausgabe is irrelevant for correction; compare the extracted month to itself.
            expected_raw = extracted.values_raw.get("ausgabe", "")
            expected_norm = extracted.values_norm.get("ausgabe", "")
            pdf_results.append(validate_one_pdf(extracted, excel_records, dok_idx, expected_raw, expected_norm))
        apply_word_page_count_checks(word_results, pdf_results, pdfs_available=True)
    else:
        apply_word_page_count_checks(word_results, pdf_results, pdfs_available=False)

    doc = Document(word_docx_path)
    if info.table_index_zero_based >= len(doc.tables):
        raise ValueError("The recognized Word table could not be reopened for correction.")
    table = doc.tables[info.table_index_zero_based]
    result = WordCorrectionResult(output_path=output_path)
    freigabe_values: set[str] = set()

    for word_result in word_results:
        if not word_result.matched_excel_row:
            result.skipped_rows += 1
            result.warnings.append(f"Word row {word_result.record.word_row_number} was skipped because it could not be matched uniquely to Excel.")
            continue
        excel_record = excel_by_row[word_result.matched_excel_row]
        row_idx = word_result.record.word_row_number - 1
        if not (0 <= row_idx < len(table.rows)):
            result.skipped_rows += 1
            result.warnings.append(f"Word row {word_result.record.word_row_number} could not be located in the output document.")
            continue
        cells = table.rows[row_idx].cells
        row_changed = False
        for field_name in ("artikelnummer", "dokumentnummer", "revision", "version"):
            col_idx = info.columns_by_index_zero_based.get(field_name)
            if col_idx is None or not (0 <= col_idx < len(cells)):
                continue
            desired = excel_record.values_raw.get(field_name, "")
            current = _word_cell_to_text(cells[col_idx])
            if current != desired:
                _set_word_cell_text_preserve_structure(cells[col_idx], desired)
                result.changed_cells += 1
                row_changed = True
        page_col = info.columns_by_index_zero_based.get("seiten")
        if page_col is not None and 0 <= page_col < len(cells) and word_result.pdf_pages is not None:
            desired_pages = str(word_result.pdf_pages)
            current_pages = _word_cell_to_text(cells[page_col])
            if current_pages != desired_pages:
                _set_word_cell_text_preserve_structure(cells[page_col], desired_pages)
                result.changed_cells += 1
                row_changed = True
        if row_changed:
            result.corrected_rows += 1
        freigabe = clean_text(excel_record.values_raw.get("freigabe", ""))
        if freigabe:
            freigabe_values.add(freigabe)

    if len(freigabe_values) == 1:
        desired_number = next(iter(freigabe_values))
        number_row = info.change_number_cell_row_zero_based
        number_col = info.change_number_cell_col_zero_based
        if info.change_number_source.startswith("section[0].header.table[0]"):
            section_idx = info.change_number_section_zero_based
            header_table_idx = info.change_number_header_table_zero_based
            if None not in {section_idx, header_table_idx, number_row, number_col}:
                try:
                    number_cell = (
                        doc.sections[section_idx]
                        .header.tables[header_table_idx]
                        .rows[number_row]
                        .cells[number_col]
                    )
                    current_paragraph = (
                        _paragraph_visible_text(number_cell.paragraphs[0])
                        if number_cell.paragraphs else _word_cell_to_text(number_cell)
                    )
                    current_first_word = current_paragraph.split()[0] if current_paragraph.split() else ""
                    if _clean_change_number_candidate(current_first_word) != desired_number:
                        _set_first_word_in_confirmed_header_cell(number_cell, desired_number)
                        result.changed_cells += 1
                except (IndexError, AttributeError) as exc:
                    result.warnings.append(
                        "The confirmed header Number address could not be written: " + str(exc)
                    )
            else:
                result.warnings.append("The confirmed header Number address is incomplete.")
        elif number_row is not None and number_col is not None and 0 <= number_row < len(table.rows):
            number_cells = table.rows[number_row].cells
            if 0 <= number_col < len(number_cells):
                number_cell = number_cells[number_col]
                if _clean_change_number_candidate(_word_cell_to_text(number_cell)) != desired_number:
                    _set_word_cell_text_preserve_structure(number_cell, desired_number)
                    result.changed_cells += 1
            else:
                result.warnings.append("The top-right Number cell could not be addressed in the output document.")
        else:
            result.warnings.append("The top-right Number cell could not be located for automatic correction.")
    elif len(freigabe_values) > 1:
        result.warnings.append(
            "The matched Excel rows contain different Freigabe-/Änd.-Nr. values ("
            + ", ".join(sorted(freigabe_values))
            + "). The single top-right Word Number was therefore left unchanged."
        )
    else:
        result.warnings.append("No reliable Excel Freigabe-/Änd.-Nr. was available for the top-right Number field.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    if progress_callback:
        progress_callback(f"Corrected Word copy saved: {output_path}")
    return result



def _append_cloned_word_row(table, template_row_index: int):
    """Append a formatting-preserving clone of an existing Word table row."""
    if not table.rows:
        return table.add_row()
    template_row_index = min(max(0, template_row_index), len(table.rows) - 1)
    new_tr = deepcopy(table.rows[template_row_index]._tr)
    table._tbl.append(new_tr)
    return table.rows[-1]


def _pdf_page_counts_by_excel_row(
    pdf_folder: Optional[Path | str],
    excel_records: Sequence[ExcelRecord],
    progress_callback=None,
) -> Tuple[Dict[int, int], List[str]]:
    """Return unambiguous actual PDF page counts keyed by matched Excel row."""
    if not clean_text(pdf_folder):
        return {}, ["No PDF folder was selected; Blatt / sheets was left blank."]
    warnings: List[str] = []
    dok_idx = build_dok_id_index(excel_records)
    matches: Dict[int, List[Tuple[str, int]]] = {}
    pdf_files = list_pdf_files(pdf_folder)
    if not pdf_files:
        return {}, ["No PDF files were found; Blatt / sheets was left blank."]
    for index, pdf in enumerate(pdf_files, start=1):
        if progress_callback:
            progress_callback(f"Reading PDF page count {index}/{len(pdf_files)}: {pdf.name}")
        extracted = extract_pdf_table(pdf, page_number=2)
        if extracted.error:
            warnings.append(f"{pdf.name}: page count could not be matched to Excel ({extracted.error}).")
            continue
        expected_raw = extracted.values_raw.get("ausgabe", "")
        expected_norm = extracted.values_norm.get("ausgabe", "")
        validation = validate_one_pdf(
            extracted, excel_records, dok_idx, expected_raw, expected_norm
        )
        if validation.matched_excel_row and extracted.total_pages is not None:
            matches.setdefault(validation.matched_excel_row, []).append(
                (pdf.name, int(extracted.total_pages))
            )
        else:
            warnings.append(f"{pdf.name}: no unique Excel row was found for its page count.")

    result: Dict[int, int] = {}
    for row_number, entries in matches.items():
        distinct = sorted({pages for _name, pages in entries})
        if len(entries) == 1 and len(distinct) == 1:
            result[row_number] = distinct[0]
        elif len(distinct) == 1:
            warnings.append(
                f"Excel row {row_number}: multiple PDFs matched; Blatt / sheets was left blank."
            )
        else:
            warnings.append(
                f"Excel row {row_number}: matched PDFs have different page counts; Blatt / sheets was left blank."
            )
    return result, warnings


def fill_word_copy_from_excel(
    word_docx_path: Path | str,
    excel_path: Path | str,
    output_path: Path | str,
    sheet_name: Optional[str] = None,
    pdf_folder: Optional[Path | str] = None,
    progress_callback=None,
) -> WordFillResult:
    """Fill a copy of the standard Word change-notice template from Excel.

    One Word data row is populated for every Excel record. Existing template rows
    are reused; formatting-preserving clones are appended when Excel contains more
    records than the template has blank rows. Only the in-scope fields are changed.
    The original Word file is never overwritten.
    """
    word_docx_path = Path(word_docx_path)
    output_path = Path(output_path)
    if word_docx_path.resolve() == output_path.resolve():
        raise ValueError("Choose a different output path; the original Word file is not overwritten.")
    if progress_callback:
        progress_callback("Loading Excel rows for Word filling...")
    excel_records, _sheet, _header, _columns = load_excel_records(excel_path, sheet_name)
    if not excel_records:
        raise ValueError("No data rows were found in Excel.")

    # Detection works for both an empty template and a previously populated file.
    _existing_records, info = extract_change_notice_word_records(word_docx_path)
    doc = Document(word_docx_path)
    if info.table_index_zero_based >= len(doc.tables):
        raise ValueError("The recognized Word table could not be reopened for filling.")
    table = doc.tables[info.table_index_zero_based]
    data_start = info.data_start_row_zero_based
    # Automatic extraction skips completely blank template rows because they are
    # irrelevant during validation. Filling must instead start at the first row
    # immediately below the detected headers.
    if data_start >= len(table.rows) and info.header_rows_zero_based:
        data_start = max(info.header_rows_zero_based) + 1
    if data_start > len(table.rows):
        raise ValueError("The first data row in the Word template could not be located.")

    pages_by_excel, page_warnings = _pdf_page_counts_by_excel_row(
        pdf_folder, excel_records, progress_callback=progress_callback
    )
    result = WordFillResult(output_path=output_path, warnings=list(page_warnings))

    # Ensure the template has enough formatted data rows.
    existing_data_rows = max(0, len(table.rows) - data_start)
    template_row_idx = data_start if data_start < len(table.rows) else len(table.rows) - 1
    while len(table.rows) - data_start < len(excel_records):
        _append_cloned_word_row(table, template_row_idx)
        result.added_rows += 1

    fields = ("artikelnummer", "dokumentnummer", "revision", "version")
    for offset, excel_record in enumerate(excel_records):
        if progress_callback:
            progress_callback(f"Filling Word row {offset + 1}/{len(excel_records)}...")
        row = table.rows[data_start + offset]
        cells = row.cells
        for field_name in fields:
            col_idx = info.columns_by_index_zero_based.get(field_name)
            if col_idx is None or not (0 <= col_idx < len(cells)):
                continue
            desired = excel_record.values_raw.get(field_name, "")
            current = _word_cell_to_text(cells[col_idx])
            if current != desired:
                _set_word_cell_text_preserve_structure(cells[col_idx], desired)
                result.changed_cells += 1
        page_col = info.columns_by_index_zero_based.get("seiten")
        if page_col is not None and 0 <= page_col < len(cells):
            desired_pages = str(pages_by_excel.get(excel_record.row_number, ""))
            current_pages = _word_cell_to_text(cells[page_col])
            if current_pages != desired_pages:
                _set_word_cell_text_preserve_structure(cells[page_col], desired_pages)
                result.changed_cells += 1
        result.filled_rows += 1

    # Clear stale in-scope values from unused rows, without touching unrelated columns.
    for row_idx in range(data_start + len(excel_records), len(table.rows)):
        cells = table.rows[row_idx].cells
        row_changed = False
        for field_name in (*fields, "seiten"):
            col_idx = info.columns_by_index_zero_based.get(field_name)
            if col_idx is None or not (0 <= col_idx < len(cells)):
                continue
            if _word_cell_to_text(cells[col_idx]):
                _set_word_cell_text_preserve_structure(cells[col_idx], "")
                result.changed_cells += 1
                row_changed = True
        if row_changed:
            result.cleared_rows += 1

    # The confirmed header Number is a single global value for the document.
    freigabe_values = {
        clean_text(record.values_raw.get("freigabe", ""))
        for record in excel_records
        if clean_text(record.values_raw.get("freigabe", ""))
    }
    if len(freigabe_values) == 1:
        desired_number = next(iter(freigabe_values))
        try:
            number_cell = doc.sections[0].header.tables[0].rows[0].cells[1]
            current_text = _paragraph_visible_text(number_cell.paragraphs[0]) if number_cell.paragraphs else _word_cell_to_text(number_cell)
            current_first = current_text.split()[0] if current_text.split() else ""
            if _clean_change_number_candidate(current_first) != desired_number:
                _set_first_word_in_confirmed_header_cell(number_cell, desired_number)
                result.changed_cells += 1
        except (IndexError, AttributeError) as exc:
            result.warnings.append("The confirmed header Number address could not be written: " + str(exc))
    elif len(freigabe_values) > 1:
        result.warnings.append(
            "Excel contains different Freigabe-/Änd.-Nr. values ("
            + ", ".join(sorted(freigabe_values))
            + "); the single Word Number was left unchanged."
        )
    else:
        result.warnings.append("Excel contains no Freigabe-/Änd.-Nr.; the Word Number was left unchanged.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    if progress_callback:
        progress_callback(f"Filled Word copy saved: {output_path}")
    return result


def _status_fill_for_severity(severity: str) -> PatternFill:
    color = {"OK": "DDF3E8", "WARNING": "FFF1CF", "ERROR": "FADDE1"}.get(severity, "FFFFFF")
    return PatternFill("solid", fgColor=color)


def _style_report_table(ws, header_row: int, max_widths: Optional[Dict[int, float]] = None) -> None:
    max_widths = max_widths or {}
    header_fill = PatternFill("solid", fgColor="007F91")
    header_font = Font(color="FFFFFF", bold=True)
    border_side = Side(style="thin", color="D7E2E7")
    border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    for row in ws.iter_rows(min_row=header_row, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = f"A{header_row + 1}"
    if ws.max_row >= header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 8
        for cell in ws[letter]:
            value = clean_text(cell.value)
            if value:
                max_len = max(max_len, min(len(value), 80))
        width = min(max_len + 2, max_widths.get(col, 36))
        ws.column_dimensions[letter].width = width


def _style_report_title(ws, title: str, subtitle: str, last_column: int) -> None:
    end = get_column_letter(last_column)
    ws.merge_cells(f"A1:{end}1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Aptos Display", size=18, bold=True, color="20343C")
    ws["A1"].fill = PatternFill("solid", fgColor="EAF5F7")
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f"A2:{end}2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=9, color="5D7078")
    ws["A2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 22


def _comparison_result_label(value: str) -> str:
    return {
        "OK": "OK",
        "BOTH_EMPTY": "OK — both empty",
        "MISMATCH": "Mismatch",
        "PDF_MISSING": "Missing in PDF",
        "REFERENCE_MISSING": "Missing in reference",
        "WORD_MISSING": "Missing in Word",
        "EXCEL_MISSING": "Missing in Excel",
        "PDF_UNSUPPORTED_DATE_FORMAT": "Unsupported PDF date format",
    }.get(value, value.replace("_", " ").title())


def write_report(
    results: Sequence[ValidationResult],
    excel_records: Sequence[ExcelRecord],
    output_path: Path | str,
    expected_ausgabe: str,
    expected_ausgabe_norm: str,
    sheet_name: str,
    word_results: Optional[Sequence[WordValidationResult]] = None,
    word_records: Optional[Sequence[WordRecord]] = None,
    word_mapping: Optional[WordMapping] = None,
    word_extraction_info: Optional[WordExtractionInfo] = None,
    word_docx_path: Optional[Path | str] = None,
    pdf_folder_provided: bool = True,
    excel_path: Optional[Path | str] = None,
    pdf_folder: Optional[Path | str] = None,
):
    """Create a compact two-sheet report: Overview and Issues.

    Raw document values are shown without generic normalization. Language, Ausgabe
    and page counts use their explicit business conversion rules only.
    """
    output_path = Path(output_path)
    word_results_list = list(word_results or [])
    results_list = list(results)

    pdf_by_excel: Dict[int, List[ValidationResult]] = {}
    unmatched_pdfs: List[ValidationResult] = []
    for result in results_list:
        if result.matched_excel_row:
            pdf_by_excel.setdefault(result.matched_excel_row, []).append(result)
        else:
            unmatched_pdfs.append(result)

    word_by_excel: Dict[int, List[WordValidationResult]] = {}
    unmatched_word: List[WordValidationResult] = []
    for result in word_results_list:
        if result.matched_excel_row:
            word_by_excel.setdefault(result.matched_excel_row, []).append(result)
        else:
            unmatched_word.append(result)

    overview_rows: List[List[object]] = []
    overview_severities: List[str] = []
    issue_rows: List[List[object]] = []

    def add_issue(
        severity: str,
        document: str,
        excel_row: object,
        word_row: object,
        pdf_file: str,
        comparison: str,
        field: str,
        source_a: str,
        value_a: object,
        source_b: str,
        value_b: object,
        details: str,
    ) -> None:
        issue_rows.append([
            severity, document, excel_row, word_row, pdf_file, comparison, field,
            source_a, clean_text(value_a), source_b, clean_text(value_b), details,
        ])

    if word_results_list and not pdf_folder_provided:
        add_issue(
            "WARNING", "Run limitation", "", "", "", "Page-count verification", "Blatt / sheets",
            "Word", "Values available", "PDF folder", "Not provided",
            "No PDF folder was selected, so the true number of pages could not be verified.",
        )

    # Detailed PDF issues.
    pdf_matching_statuses = {
        "LIKELY_WRONG_DOK_ID", "DUPLICATE_DOK_ID_RESOLVED", "DUPLICATE_DOK_ID_AMBIGUOUS",
        "DOK_ID_NOT_FOUND", "NO_RELIABLE_MATCH", "EXTRACTION_ERROR",
    }
    for result in results_list:
        ext = result.extracted
        document = ext.values_raw.get("dokumentnummer", "") or ext.file_name
        if result.status in pdf_matching_statuses:
            add_issue(
                _pdf_severity(result.status), document, result.matched_excel_row or "", "", ext.file_name,
                "PDF matching", "Match", "PDF", ext.values_raw.get("dok_id", ""), "Excel", result.matched_excel_row or "Not found",
                _join_unique(result.issues) or PDF_STATUS_LABELS.get(result.status, result.status),
            )
        for warning in ext.warnings:
            add_issue("WARNING", document, result.matched_excel_row or "", "", ext.file_name,
                      "PDF extraction", "Extraction", "PDF", "", "", "", warning)
        for comp in result.comparisons:
            if comp.result in {"OK", "BOTH_EMPTY"}:
                continue
            severity = "WARNING" if result.status in {"LIKELY_WRONG_DOK_ID", "DUPLICATE_DOK_ID_RESOLVED"} and comp.field == PDF_FIELD_DISPLAY["dok_id"] else "ERROR"
            add_issue(
                severity, document, comp.excel_row or result.matched_excel_row or "", "", ext.file_name,
                "PDF vs " + comp.source, comp.field, "PDF", comp.pdf_raw, comp.source, comp.reference_raw,
                _comparison_result_label(comp.result),
            )

    # Detailed Word issues.
    word_matching_statuses = {
        "WORD_DUPLICATE_DOCUMENT_NUMBER_RESOLVED", "WORD_LIKELY_WRONG_DOCUMENT_NUMBER",
        "WORD_LIKELY_WRONG_DOK_ID", "WORD_DOCUMENT_NOT_IN_EXCEL", "WORD_AMBIGUOUS_MATCH",
        "WORD_NO_RELIABLE_MATCH", "WORD_TABLE_NOT_RECOGNIZED", "WORD_EXTRACTION_ERROR",
    }
    for result in word_results_list:
        record = result.record
        document = record.values_raw.get("dokumentnummer", "") or f"Word row {record.word_row_number}"
        if result.status in word_matching_statuses:
            add_issue(
                _word_severity(result), document, result.matched_excel_row or "", record.word_row_number, result.pdf_file_name,
                "Word matching", "Match", "Word", record.values_raw.get("dokumentnummer", ""), "Excel", result.matched_excel_row or "Not found",
                _join_unique(result.issues) or WORD_STATUS_LABELS.get(result.status, result.status),
            )
        for comp in result.comparisons:
            if comp.result in {"OK", "BOTH_EMPTY"}:
                continue
            severity = "WARNING" if result.status in {
                "WORD_LIKELY_WRONG_DOCUMENT_NUMBER", "WORD_LIKELY_WRONG_DOK_ID",
                "WORD_DUPLICATE_DOCUMENT_NUMBER_RESOLVED",
            } and comp.field in {PDF_FIELD_DISPLAY["dokumentnummer"], PDF_FIELD_DISPLAY["dok_id"]} else "ERROR"
            add_issue(
                severity, document, comp.excel_row or result.matched_excel_row or "", comp.word_row_number, result.pdf_file_name,
                "Word vs Excel", comp.field, "Word", comp.word_raw, "Excel", comp.excel_raw,
                _comparison_result_label(comp.result),
            )
        if result.page_count_result not in {"OK", "NOT_CHECKED", "", "UNAVAILABLE_NO_PDFS"}:
            add_issue(
                "ERROR", document, result.matched_excel_row or "", record.word_row_number, result.pdf_file_name,
                "Word vs PDF", "Blatt / sheets", "Word", result.word_pages_raw, "PDF", result.pdf_pages if result.pdf_pages is not None else "Unavailable",
                result.page_count_issue or PAGE_RESULT_LABELS.get(result.page_count_result, result.page_count_result),
            )

    # One Overview row for each Excel record.
    for record in excel_records:
        pdf_matches = pdf_by_excel.get(record.row_number, [])
        word_matches = word_by_excel.get(record.row_number, [])
        source_severities: List[str] = []
        row_issues: List[str] = []

        if pdf_folder_provided:
            if pdf_matches:
                source_severities.extend(_pdf_severity(r.status) for r in pdf_matches)
                for pdf_result in pdf_matches:
                    if pdf_result.status not in {"OK", "OK_WITH_NORMALIZATION"}:
                        row_issues.append(PDF_STATUS_LABELS.get(pdf_result.status, pdf_result.status))
                    row_issues.extend(
                        comp.field for comp in pdf_result.comparisons
                        if comp.result not in {"OK", "BOTH_EMPTY"}
                    )
                    if pdf_result.extracted.warnings:
                        row_issues.append("PDF extraction warning")
            else:
                source_severities.append("ERROR")
                missing = "Excel row was not matched to any PDF in the selected folder."
                row_issues.append(missing)
                add_issue(
                    "ERROR", record.values_raw.get("dokumentnummer", "") or record.values_raw.get("dok_id", ""), record.row_number, "", "",
                    "PDF presence", "Document", "Excel", record.values_raw.get("dokumentnummer", ""), "PDF folder", "Not found", missing,
                )
        if word_results is not None:
            if word_matches:
                source_severities.extend(_word_severity(r) for r in word_matches)
                for word_result in word_matches:
                    if word_result.status not in {"WORD_OK", "WORD_OK_WITH_NORMALIZATION"}:
                        row_issues.append(WORD_STATUS_LABELS.get(word_result.status, word_result.status))
                    row_issues.extend(
                        comp.field for comp in word_result.comparisons
                        if comp.result not in {"OK", "BOTH_EMPTY"}
                    )
                    if word_result.page_count_result not in {"OK", "NOT_CHECKED", ""}:
                        row_issues.append(PAGE_RESULT_LABELS.get(word_result.page_count_result, word_result.page_count_result))
            else:
                source_severities.append("ERROR")
                missing = "Excel row was not found in the selected Word table."
                row_issues.append(missing)
                add_issue(
                    "ERROR", record.values_raw.get("dokumentnummer", "") or record.values_raw.get("dok_id", ""), record.row_number, "", "",
                    "Word presence", "Document", "Excel", record.values_raw.get("dokumentnummer", ""), "Word table", "Not found", missing,
                )

        severity = _max_severity(source_severities or ["OK"])
        pdf_status = _join_unique(PDF_STATUS_LABELS.get(r.status, r.status) for r in pdf_matches) if pdf_matches else ("Not found" if pdf_folder_provided else "Not provided")
        word_status = _join_unique(WORD_STATUS_LABELS.get(r.excel_status or r.status, r.excel_status or r.status) for r in word_matches) if word_matches else ("Not found" if word_results is not None else "Not provided")
        ausgabe_comps = [c for r in pdf_matches for c in r.comparisons if c.field == PDF_FIELD_DISPLAY["ausgabe"]]
        ausgabe_check = _join_unique(_comparison_result_label(c.result) for c in ausgabe_comps) if ausgabe_comps else ("Not available" if not pdf_folder_provided else "Not found")
        page_check = _join_unique(PAGE_RESULT_LABELS.get(r.page_count_result, r.page_count_result) for r in word_matches) if word_matches else ("Not applicable" if word_results is None else "Not found")
        overview_rows.append([
            severity,
            record.row_number,
            record.values_raw.get("dok_id", ""),
            record.values_raw.get("dokumentnummer", ""),
            record.values_raw.get("artikelnummer", ""),
            _join_unique(r.extracted.file_name for r in pdf_matches),
            _join_unique(r.record.word_row_number for r in word_matches),
            pdf_status,
            word_status,
            _join_unique(c.pdf_raw for c in ausgabe_comps),
            expected_ausgabe if pdf_folder_provided else "",
            ausgabe_check,
            _join_unique(r.word_pages_raw for r in word_matches),
            _join_unique(r.pdf_pages for r in word_matches if r.pdf_pages is not None) or _join_unique(r.extracted.total_pages for r in pdf_matches if r.extracted.total_pages is not None),
            page_check,
            _join_unique(row_issues),
        ])
        overview_severities.append(severity)

    # Source records that could not be linked to an Excel row remain visible.
    for result in unmatched_pdfs:
        ext = result.extracted
        severity = _pdf_severity(result.status)
        overview_rows.append([
            severity, "", ext.values_raw.get("dok_id", ""), ext.values_raw.get("dokumentnummer", ""),
            ext.values_raw.get("artikelnummer", ""), ext.file_name, "",
            PDF_STATUS_LABELS.get(result.status, result.status), "Not linked", ext.values_raw.get("ausgabe", ""),
            expected_ausgabe, _comparison_result_label(next((c.result for c in result.comparisons if c.field == PDF_FIELD_DISPLAY["ausgabe"]), "")),
            "", ext.total_pages if ext.total_pages is not None else "", "Not linked", _join_unique(result.issues + ext.warnings),
        ])
        overview_severities.append(severity)

    for result in unmatched_word:
        rec = result.record
        severity = _word_severity(result)
        overview_rows.append([
            severity, "", rec.values_raw.get("dok_id", ""), rec.values_raw.get("dokumentnummer", ""),
            rec.values_raw.get("artikelnummer", ""), result.pdf_file_name, rec.word_row_number,
            "Not linked", WORD_STATUS_LABELS.get(result.status, result.status), "", "", "Not applicable",
            result.word_pages_raw, result.pdf_pages if result.pdf_pages is not None else "",
            PAGE_RESULT_LABELS.get(result.page_count_result, result.page_count_result), _join_unique(result.issues),
        ])
        overview_severities.append(severity)

    counts = {"OK": 0, "WARNING": 0, "ERROR": 0}
    for severity in overview_severities:
        counts[severity] = counts.get(severity, 0) + 1

    wb = Workbook()
    overview = wb.active
    overview.title = "Overview"
    overview_headers = [
        "Status", "Excel row", "DOK-ID", "Dok.-Nr.", "Artikel-Nr.", "PDF file(s)", "Word row(s)",
        "PDF ↔ Excel", "Word ↔ Excel", "PDF Ausgabe", "Expected Ausgabe", "Ausgabe check",
        "Word Blatt / sheets", "Actual PDF pages", "Page check", "Issues",
    ]
    _style_report_title(
        overview,
        "Profile Docs Checker — Validation Overview",
        "Compact document-level summary. Exact values are preserved; only language, Ausgabe and page count use defined conversion rules.",
        len(overview_headers),
    )
    metadata = [
        ("Excel file", Path(excel_path).name if excel_path else ""),
        ("Excel sheet", sheet_name),
        ("PDF folder", Path(pdf_folder).name if pdf_folder_provided and pdf_folder else ("Not provided" if not pdf_folder_provided else "Provided")),
        ("Word file", Path(word_docx_path).name if word_docx_path else "Not provided"),
        ("Expected Ausgabe", expected_ausgabe if pdf_folder_provided else "Not applicable"),
    ]
    for idx, (label, value) in enumerate(metadata, start=3):
        overview.cell(idx, 1, label).font = Font(bold=True, color="5D7078")
        overview.cell(idx, 2, value)
    metrics = [
        ("Excel records", len(excel_records)),
        ("PDFs processed", len(results_list)),
        ("Word rows", len(word_results_list)),
        ("Issue entries", len(issue_rows)),
    ]
    for idx, (label, value) in enumerate(metrics, start=3):
        overview.cell(idx, 4, label).font = Font(bold=True, color="5D7078")
        overview.cell(idx, 5, value)
    for idx, severity in enumerate(["OK", "WARNING", "ERROR"], start=3):
        overview.cell(idx, 7, severity).font = Font(bold=True, color="5D7078")
        overview.cell(idx, 8, counts.get(severity, 0))
        overview.cell(idx, 8).fill = _status_fill_for_severity(severity)
    header_row = 9
    for col_idx, header in enumerate(overview_headers, start=1):
        overview.cell(header_row, col_idx, header)
    for row, severity in zip(overview_rows, overview_severities):
        overview.append(row)
        overview.cell(overview.max_row, 1).fill = _status_fill_for_severity(severity)
        overview.cell(overview.max_row, 1).font = Font(bold=True)
    _style_report_table(
        overview,
        header_row,
        max_widths={1: 12, 2: 11, 3: 20, 4: 22, 5: 22, 6: 32, 7: 13, 8: 28, 9: 30, 10: 18, 11: 18, 12: 22, 13: 18, 14: 17, 15: 28, 16: 55},
    )

    issues = wb.create_sheet("Issues")
    issue_headers = [
        "Severity", "Document", "Excel row", "Word row", "PDF file", "Comparison", "Field",
        "Source A", "Value A", "Source B", "Value B", "Details",
    ]
    _style_report_title(
        issues,
        "Profile Docs Checker — Issues",
        "Only warnings, mismatches, missing records and run limitations are listed here.",
        len(issue_headers),
    )
    issue_header_row = 4
    for col_idx, header in enumerate(issue_headers, start=1):
        issues.cell(issue_header_row, col_idx, header)
    if issue_rows:
        for row in sorted(issue_rows, key=lambda r: (-_severity_rank(str(r[0])), clean_text(r[1]), clean_text(r[5]), clean_text(r[6]))):
            issues.append(row)
            issues.cell(issues.max_row, 1).fill = _status_fill_for_severity(str(row[0]))
            issues.cell(issues.max_row, 1).font = Font(bold=True)
    else:
        issues.append(["OK", "No issues found", "", "", "", "Validation", "", "", "", "", "", "All selected checks passed."])
        issues.cell(issues.max_row, 1).fill = _status_fill_for_severity("OK")
    _style_report_table(
        issues,
        issue_header_row,
        max_widths={1: 12, 2: 24, 3: 11, 4: 11, 5: 30, 6: 22, 7: 24, 8: 16, 9: 30, 10: 16, 11: 30, 12: 60},
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

def summarize_results(results: Sequence[ValidationResult]) -> Dict[str, int]:
    counts: Dict[str, int] = {}
    for res in results:
        counts[res.status] = counts.get(res.status, 0) + 1
    return dict(sorted(counts.items(), key=lambda kv: STATUS_ORDER.get(kv[0], 99)))
