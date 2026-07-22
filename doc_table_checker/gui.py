from __future__ import annotations

import os
import threading
import webbrowser
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook

from .core import (
    create_corrected_word_copy,
    fill_word_copy_from_excel,
    run_validation,
)


TRANSLATIONS = {
    "de": {
        "inputs": "PRÜFEINGABEN",
        "control": "STEUERUNG",
        "excel_required": "EXCEL-MASTERLISTE  •  ERFORDERLICH",
        "excel_hint": "Feste deutsche Mastervorlage (.xlsx / .xlsm)",
        "sheet": "EXCEL-TABELLE",
        "reload": "NEU LADEN",
        "word_optional": "WORD-ÄNDERUNGSMITTEILUNG  •  OPTIONAL",
        "word_hint": "Word-Datei auswählen und anschließend zwischen Prüfen und Befüllen wählen.",
        "word_action": "WORD-AKTION",
        "examine": "PRÜFEN",
        "fill": "AUS EXCEL BEFÜLLEN",
        "pdf_optional": "PDF-DOKUMENTORDNER  •  OPTIONAL",
        "pdf_hint": "Für PDF-Metadaten und echte Seitenzahlen. Metadaten werden aus Seite 2 gelesen.",
        "expected_issue": "ERWARTETE AUSGABE",
        "issue_hint": "Monat und Jahr, z. B. 07.2026. Nur bei PDF-Prüfung erforderlich.",
        "output_report": "AUSGABEBERICHT",
        "output_report_hint": "Kompakter Excel-Prüfbericht mit zwei Tabellenblättern (.xlsx)",
        "output_word": "AUSGABE-WORD-DATEI",
        "output_word_hint": "Neue befüllte Word-Kopie; die Originaldatei bleibt unverändert (.docx)",
        "browse": "AUSWÄHLEN",
        "clear": "LEEREN",
        "save_as": "SPEICHERN UNTER",
        "run_check": "DOKUMENTE PRÜFEN UND BERICHT ERSTELLEN",
        "run_fill": "WORD-DATEI AUS EXCEL BEFÜLLEN",
        "open_report": "EXCEL-BERICHT ÖFFNEN",
        "result": "PRÜFERGEBNIS",
        "idle": "BEREIT",
        "mode": "MODUS",
        "id_heading": "ID-NR.",
        "issue_heading": "FALSCHE ODER NICHT PRÜFBARE ANGABE",
        "correct_copy": "KORRIGIERTE WORD-KOPIE ERSTELLEN",
        "await_title": "BEREIT ZUR PRÜFUNG",
        "await_detail": "Das kompakte Gesamtergebnis erscheint nach der Prüfung hier.",
        "no_result": "Noch kein Prüfergebnis.",
        "select_sources": "Excel-Datei und mindestens eine Dokumentquelle auswählen.",
        "mode_full": "VOLLPRÜFUNG  •  PDF ↔ Excel  •  Word ↔ Excel  •  Word-Seiten ↔ tatsächliche PDF-Seiten",
        "mode_pdf": "PDF-PRÜFUNG  •  PDF-Metadaten ↔ Excel  •  Ausgabe-Prüfung",
        "mode_word": "WORD-PRÜFUNG  •  Word ↔ Excel  •  ohne PDFs bleiben Seitenzahlen ungeprüft",
        "mode_fill_pdf": "WORD BEFÜLLEN  •  Werte aus Excel  •  Blatt/Seiten aus PDFs",
        "mode_fill_no_pdf": "WORD BEFÜLLEN  •  Werte aus Excel  •  Blatt/Seiten bleiben ohne PDFs leer",
        "running": "LÄUFT",
        "complete": "FERTIG",
        "error": "FEHLER",
        "validation_progress": "PRÜFUNG LÄUFT",
        "validation_detail": "Das kompakte Ergebnis erscheint, sobald alle gewählten Prüfungen abgeschlossen sind.",
        "checking": "Dokumente werden geprüft…",
        "start_validation": "Prüfung wird gestartet…",
        "all_correct": "ALLE EINTRÄGE SIND KORREKT",
        "all_correct_detail": "Alle gewählten Prüfungen wurden für {count} Excel-{entry} bestanden.",
        "entry_singular": "Eintrag",
        "entry_plural": "Einträge",
        "all_values_correct": "Alle gewählten Werte sind korrekt.",
        "items_attention": "{count} {items} ZU PRÜFEN",
        "item_singular": "ANGABE",
        "item_plural": "ANGABEN",
        "errors_warnings": "{errors} {error_word}{warnings_text}. Falsche oder nicht prüfbare Werte sind unten markiert.",
        "error_singular": "Fehler",
        "error_plural": "Fehler",
        "warning_singular": "Warnung",
        "warning_plural": "Warnungen",
        "validation_failed": "PRÜFUNG FEHLGESCHLAGEN",
        "validation_did_not_finish": "Die Prüfung wurde nicht abgeschlossen.",
        "missing_input": "Fehlende oder ungültige Eingabe",
        "choose_excel": "Excel-Masterliste auswählen.",
        "choose_source": "PDF-Ordner, Word-Datei oder beides auswählen.",
        "choose_word_fill": "Für das Befüllen eine Word-Vorlage auswählen.",
        "enter_issue": "Bei ausgewähltem PDF-Ordner die erwartete Ausgabe eingeben.",
        "choose_output": "Ausgabepfad auswählen.",
        "page_unavailable_title": "Tatsächliche Seitenzahl nicht verfügbar",
        "page_unavailable_check": "Es wurde kein PDF-Ordner ausgewählt.\n\nDie Word-Tabelle wird mit Excel verglichen, aber Blatt/Seiten kann nicht gegen die tatsächliche PDF-Seitenzahl geprüft werden. Fortfahren?",
        "page_unavailable_fill": "Es wurde kein PDF-Ordner ausgewählt.\n\nDie Word-Datei wird aus Excel befüllt; Blatt/Seiten bleibt jedoch leer. Fortfahren?",
        "report_saved": "Excel-Bericht gespeichert: {path}",
        "filled_title": "WORD-DATEI ERFOLGREICH BEFÜLLT",
        "filled_detail": "{rows} Excel-Zeilen wurden in eine neue Word-Kopie übernommen.",
        "filled_success": "Word-Datei wurde aus Excel befüllt.",
        "filled_warning": "Hinweis: {warning}",
        "filled_saved": "Befüllte Word-Kopie gespeichert: {path}",
        "fill_failed": "WORD-BEFÜLLUNG FEHLGESCHLAGEN",
        "fill_did_not_finish": "Die Word-Datei konnte nicht befüllt werden.",
        "no_word_result": "Kein korrigierbares Word-Ergebnis",
        "no_word_result_msg": "Zuerst eine Prüfung mit einer fehlerhaften Word-Datei durchführen.",
        "save_corrected": "Korrigierte Word-Kopie speichern",
        "create_corrected_title": "Korrigierte Word-Kopie erstellen",
        "create_corrected_question": "Es wird eine neue Word-Datei erstellt; das Original bleibt unverändert.\n\nZuverlässig zugeordnete Zeilen erhalten die exakten Excel-Werte. Blatt/Seiten verwendet die tatsächlichen PDF-Seitenzahlen, sofern PDFs vorhanden sind. Fortfahren?",
        "correcting": "KORREKTUR",
        "creating_corrected": "Korrigierte Word-Kopie wird erstellt…",
        "corrected_saved": "Korrigierte Word-Kopie gespeichert: {path}",
        "corrected_created": "Korrigierte Word-Kopie erstellt",
        "corrected_info": "Gespeichert unter:\n{path}\n\nGeänderte Zellen: {cells}\nKorrigierte Zeilen: {rows}\nÜbersprungene Zeilen: {skipped}{warnings}",
        "warnings": "Warnungen",
        "correction_failed": "Word-Korrektur fehlgeschlagen",
        "could_not_correct": "Die korrigierte Word-Kopie konnte nicht erstellt werden.",
        "report_not_found": "Bericht nicht gefunden",
        "report_not_found_msg": "Die Berichtsdatei ist noch nicht vorhanden.",
        "could_not_open_report": "Bericht konnte nicht geöffnet werden",
        "choose_pdf_folder": "Ordner mit PDF-Dateien auswählen",
        "choose_excel_file": "Excel-Masterliste auswählen",
        "choose_word_file": "Word-Änderungsmitteilung auswählen",
        "choose_report_path": "Prüfbericht speichern unter",
        "choose_filled_word_path": "Befüllte Word-Kopie speichern unter",
        "sheets_loaded": "Excel-Tabellen geladen: {sheets}",
        "sheet_load_error": "Excel-Tabellen konnten nicht geladen werden",
        "language": "SPRACHE",
    },
    "en": {
        "inputs": "VALIDATION INPUTS",
        "control": "CONTROL",
        "excel_required": "EXCEL MASTER LIST  •  REQUIRED",
        "excel_hint": "Fixed German master template (.xlsx / .xlsm)",
        "sheet": "EXCEL SHEET",
        "reload": "RELOAD",
        "word_optional": "WORD CHANGE NOTICE  •  OPTIONAL",
        "word_hint": "Choose a Word file, then select whether to examine it or fill it from Excel.",
        "word_action": "WORD ACTION",
        "examine": "EXAMINE",
        "fill": "FILL FROM EXCEL",
        "pdf_optional": "PDF DOCUMENT FOLDER  •  OPTIONAL",
        "pdf_hint": "Required for PDF metadata and actual page counts. Metadata is read from page 2.",
        "expected_issue": "EXPECTED AUSGABE",
        "issue_hint": "Month and year, e.g. 07.2026. Required only for PDF validation.",
        "output_report": "OUTPUT REPORT",
        "output_report_hint": "Compact two-sheet Excel validation report (.xlsx)",
        "output_word": "OUTPUT WORD FILE",
        "output_word_hint": "New filled Word copy; the original remains unchanged (.docx)",
        "browse": "BROWSE",
        "clear": "CLEAR",
        "save_as": "SAVE AS",
        "run_check": "CHECK DOCUMENTS AND CREATE REPORT",
        "run_fill": "FILL WORD FILE FROM EXCEL",
        "open_report": "OPEN EXCEL REPORT",
        "result": "VALIDATION RESULT",
        "idle": "IDLE",
        "mode": "MODE",
        "id_heading": "ID-NR.",
        "issue_heading": "INCORRECT OR UNVERIFIED INFORMATION",
        "correct_copy": "CREATE CORRECTED WORD COPY",
        "await_title": "AWAITING VALIDATION",
        "await_detail": "The complete concise result will appear here after validation.",
        "no_result": "No validation result yet.",
        "select_sources": "Select an Excel file and at least one document source.",
        "mode_full": "FULL CONTROL  •  PDF ↔ Excel  •  Word ↔ Excel  •  Word pages ↔ actual PDF pages",
        "mode_pdf": "PDF CONTROL  •  PDF metadata ↔ Excel  •  Ausgabe validation",
        "mode_word": "WORD CONTROL  •  Word ↔ Excel  •  page count remains unverified without PDFs",
        "mode_fill_pdf": "FILL WORD  •  values from Excel  •  Blatt/pages from PDFs",
        "mode_fill_no_pdf": "FILL WORD  •  values from Excel  •  Blatt/pages remain blank without PDFs",
        "running": "RUNNING",
        "complete": "COMPLETE",
        "error": "ERROR",
        "validation_progress": "VALIDATION IN PROGRESS",
        "validation_detail": "The concise result will appear as soon as all selected checks finish.",
        "checking": "Checking documents…",
        "start_validation": "Starting validation sequence…",
        "all_correct": "ALL ENTRIES ARE CORRECT",
        "all_correct_detail": "All selected checks passed for {count} Excel {entry}.",
        "entry_singular": "entry",
        "entry_plural": "entries",
        "all_values_correct": "All selected values are correct.",
        "items_attention": "{count} {items} NEED ATTENTION",
        "item_singular": "ITEM",
        "item_plural": "ITEMS",
        "errors_warnings": "{errors} {error_word}{warnings_text}. Incorrect or unverified values are highlighted below.",
        "error_singular": "error",
        "error_plural": "errors",
        "warning_singular": "warning",
        "warning_plural": "warnings",
        "validation_failed": "VALIDATION FAILED",
        "validation_did_not_finish": "The validation did not finish.",
        "missing_input": "Missing or invalid input",
        "choose_excel": "Choose the Excel master list.",
        "choose_source": "Choose a PDF folder, a Word file, or both.",
        "choose_word_fill": "Choose a Word template to fill.",
        "enter_issue": "Enter the expected Ausgabe when a PDF folder is selected.",
        "choose_output": "Choose an output path.",
        "page_unavailable_title": "Actual page count unavailable",
        "page_unavailable_check": "No PDF folder was selected.\n\nThe Word table will be compared with Excel, but Blatt / sheets cannot be verified against the actual PDF page count. Continue?",
        "page_unavailable_fill": "No PDF folder was selected.\n\nThe Word file will be filled from Excel, but Blatt / sheets will remain blank. Continue?",
        "report_saved": "Excel report saved: {path}",
        "filled_title": "WORD FILE FILLED SUCCESSFULLY",
        "filled_detail": "{rows} Excel rows were transferred into a new Word copy.",
        "filled_success": "The Word file was filled from Excel.",
        "filled_warning": "Notice: {warning}",
        "filled_saved": "Filled Word copy saved: {path}",
        "fill_failed": "WORD FILL FAILED",
        "fill_did_not_finish": "The Word file could not be filled.",
        "no_word_result": "No correctable Word result",
        "no_word_result_msg": "First run validation with a Word file containing correctable differences.",
        "save_corrected": "Save corrected Word copy",
        "create_corrected_title": "Create corrected Word copy",
        "create_corrected_question": "A new Word file will be created; the original remains unchanged.\n\nReliably matched rows will receive exact Excel values. Blatt / sheets will use actual PDF page counts when PDFs are available. Continue?",
        "correcting": "CORRECTING",
        "creating_corrected": "Creating corrected Word copy…",
        "corrected_saved": "Corrected Word copy saved: {path}",
        "corrected_created": "Corrected Word copy created",
        "corrected_info": "Saved to:\n{path}\n\nChanged cells: {cells}\nCorrected rows: {rows}\nSkipped rows: {skipped}{warnings}",
        "warnings": "Warnings",
        "correction_failed": "Word correction failed",
        "could_not_correct": "Could not create the corrected Word copy.",
        "report_not_found": "Report not found",
        "report_not_found_msg": "The report file does not exist yet.",
        "could_not_open_report": "Could not open report",
        "choose_pdf_folder": "Choose folder containing PDFs",
        "choose_excel_file": "Choose Excel master list",
        "choose_word_file": "Choose Word change notice",
        "choose_report_path": "Save validation report as",
        "choose_filled_word_path": "Save filled Word copy as",
        "sheets_loaded": "Excel sheets loaded: {sheets}",
        "sheet_load_error": "Could not load Excel sheets",
        "language": "LANGUAGE",
    },
}


class App(tk.Tk):
    """Light control-panel interface with bilingual validation and Word filling."""

    BG = "#EEF4F6"
    PANEL = "#FFFFFF"
    PANEL_ALT = "#E8F1F4"
    PANEL_HOVER = "#DCECEF"
    BORDER = "#BFD1D8"
    ACCENT = "#00AFC3"
    ACCENT_DARK = "#007A8A"
    TEXT = "#20343C"
    MUTED = "#667A83"
    WARNING = "#AE7200"
    WARNING_BG = "#FFF4D8"
    ERROR = "#C94855"
    ERROR_BG = "#FCE4E7"
    SUCCESS = "#247F5C"
    SUCCESS_BG = "#E2F4EB"

    def __init__(self):
        super().__init__()
        self.title("Profile Docs Checker")
        self.geometry("1280x820")
        self.minsize(1100, 720)
        self.configure(bg=self.BG)

        self.pdf_folder_var = tk.StringVar()
        self.excel_file_var = tk.StringVar()
        self.sheet_var = tk.StringVar()
        self.word_file_var = tk.StringVar()
        self.word_mode_var = tk.StringVar(value="examine")
        self.ausgabe_var = tk.StringVar()
        self.output_file_var = tk.StringVar()
        self.language_var = tk.StringVar(value="Deutsch")
        self.lang = "de"

        self.mode_text_var = tk.StringVar()
        self.current_step_var = tk.StringVar()
        self.result_title_var = tk.StringVar()
        self.result_detail_var = tk.StringVar()
        self._last_summary: dict = {}
        self._last_fill_result = None
        self._last_run_succeeded = False
        self._is_running = False

        self._configure_styles()
        self._build_ui()
        self._set_initial_result()
        self._refresh_mode_status()
        self.after_idle(self._maximize_or_fit)

        self.pdf_folder_var.trace_add("write", lambda *_: self._refresh_mode_status())
        self.word_file_var.trace_add("write", lambda *_: self._refresh_mode_status())
        self.word_mode_var.trace_add("write", lambda *_: self._on_word_mode_change())

    def _t(self, key: str) -> str:
        return TRANSLATIONS[self.lang][key]

    def _maximize_or_fit(self):
        try:
            if os.name == "nt":
                self.state("zoomed")
                return
        except tk.TclError:
            pass
        self.update_idletasks()
        screen_w = self.winfo_screenwidth()
        screen_h = self.winfo_screenheight()
        width = min(1360, max(1100, screen_w - 40))
        height = min(900, max(720, screen_h - 60))
        self.geometry(f"{width}x{height}+{max(0, (screen_w-width)//2)}+{max(0, (screen_h-height)//2)}")

    def _configure_styles(self):
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("Accent.TButton", background=self.ACCENT, foreground="#FFFFFF", bordercolor=self.ACCENT,
                        lightcolor=self.ACCENT, darkcolor=self.ACCENT_DARK, relief="flat", padding=(18, 10),
                        font=("Segoe UI Semibold", 10))
        style.map("Accent.TButton", background=[("disabled", "#D7E3E7"), ("active", "#18C0D1")],
                  foreground=[("disabled", self.MUTED), ("active", "#FFFFFF")])
        style.configure("Ghost.TButton", background=self.PANEL_ALT, foreground=self.TEXT, bordercolor=self.BORDER,
                        lightcolor=self.PANEL_ALT, darkcolor=self.PANEL_ALT, relief="flat", padding=(12, 8),
                        font=("Segoe UI", 9))
        style.map("Ghost.TButton", background=[("active", self.PANEL_HOVER), ("disabled", self.PANEL)],
                  foreground=[("disabled", "#9AAAB0")])
        style.configure("Compact.TButton", background=self.PANEL_ALT, foreground=self.TEXT, bordercolor=self.BORDER,
                        lightcolor=self.PANEL_ALT, darkcolor=self.PANEL_ALT, relief="flat", padding=(8, 7),
                        font=("Segoe UI", 8))
        style.map("Compact.TButton", background=[("active", self.PANEL_HOVER)])
        style.configure("Mode.TRadiobutton", background=self.PANEL, foreground=self.TEXT, font=("Segoe UI Semibold", 9),
                        indicatorcolor="#F8FBFC", indicatorrelief="flat", padding=(6, 5))
        style.map("Mode.TRadiobutton", indicatorcolor=[("selected", self.ACCENT), ("!selected", "#F8FBFC")],
                  foreground=[("selected", self.ACCENT_DARK)])
        style.configure("Light.TCombobox", fieldbackground="#F8FBFC", background="#F8FBFC", foreground=self.TEXT,
                        arrowcolor=self.ACCENT, bordercolor=self.BORDER, lightcolor=self.BORDER,
                        darkcolor=self.BORDER, padding=5)
        style.map("Light.TCombobox", fieldbackground=[("readonly", "#F8FBFC")], foreground=[("readonly", self.TEXT)],
                  selectbackground=[("readonly", "#F8FBFC")], selectforeground=[("readonly", self.TEXT)])
        style.configure("Control.Horizontal.TProgressbar", troughcolor=self.PANEL, background=self.ACCENT,
                        bordercolor=self.BORDER, lightcolor=self.ACCENT, darkcolor=self.ACCENT_DARK, thickness=5)
        style.configure("Results.Treeview", background="#F8FBFC", fieldbackground="#F8FBFC", foreground=self.TEXT,
                        rowheight=38, bordercolor=self.BORDER, lightcolor=self.BORDER, darkcolor=self.BORDER,
                        font=("Segoe UI", 9))
        style.configure("Results.Treeview.Heading", background=self.PANEL_ALT, foreground=self.TEXT, relief="flat",
                        font=("Segoe UI Semibold", 9), padding=(8, 7))
        style.map("Results.Treeview", background=[("selected", "#CDECF1")], foreground=[("selected", self.TEXT)])

    def _build_ui(self):
        self._build_header()
        body = tk.Frame(self, bg=self.BG)
        body.pack(fill="both", expand=True, padx=24, pady=(18, 20))
        body.grid_columnconfigure(0, weight=6, minsize=540)
        body.grid_columnconfigure(1, weight=6, minsize=480)
        body.grid_rowconfigure(0, weight=1)
        left = tk.Frame(body, bg=self.PANEL, highlightthickness=1, highlightbackground=self.BORDER)
        left.grid(row=0, column=0, sticky="nsew", padx=(0, 12))
        left.grid_columnconfigure(0, weight=1)
        right = tk.Frame(body, bg=self.PANEL, highlightthickness=1, highlightbackground=self.BORDER)
        right.grid(row=0, column=1, sticky="nsew", padx=(12, 0))
        right.grid_columnconfigure(0, weight=1)
        right.grid_rowconfigure(4, weight=1)
        self._build_configuration(left)
        self._build_results_panel(right)

    def _build_header(self):
        self.header_canvas = tk.Canvas(self, height=112, bg=self.BG, highlightthickness=0)
        self.header_canvas.pack(fill="x")
        self.header_canvas.bind("<Configure>", lambda event: self._draw_header(event.width))
        language_frame = tk.Frame(self, bg=self.PANEL)
        language_frame.place(relx=1.0, x=-42, y=18, anchor="ne")
        tk.Label(language_frame, text=self._t("language"), bg=self.PANEL, fg=self.MUTED,
                 font=("Segoe UI", 8, "bold")).pack(side="left", padx=(0, 7))
        combo = ttk.Combobox(language_frame, textvariable=self.language_var, state="readonly", width=10,
                             values=("Deutsch", "English"), style="Light.TCombobox")
        combo.pack(side="left")
        combo.bind("<<ComboboxSelected>>", self._change_language)

    def _draw_header(self, width: int):
        c = self.header_canvas
        c.delete("all")
        c.create_polygon(0, 0, width, 0, width, 84, width - 38, 112, 38, 112, 0, 84,
                         fill=self.PANEL, outline=self.BORDER)
        c.create_line(42, 111, width - 42, 111, fill=self.ACCENT, width=2)
        c.create_polygon(0, 0, 210, 0, 180, 12, 0, 12, fill=self.ACCENT_DARK, outline="")
        c.create_text(44, 43, anchor="w", text="PROFILE // DOCS CHECKER", fill=self.TEXT,
                      font=("Segoe UI Semibold", 19))
        c.create_text(45, 73, anchor="w", text="DOCUMENT CONSISTENCY CONTROL", fill=self.ACCENT,
                      font=("Consolas", 9, "bold"))
        c.create_text(width - 265, 74, anchor="e", text="PDF  •  EXCEL  •  WORD", fill=self.MUTED,
                      font=("Segoe UI", 10))
        c.create_text(width - 45, 74, anchor="e", text="v1.6.0", fill=self.ACCENT, font=("Consolas", 9))

    def _build_configuration(self, parent: tk.Frame):
        title_row = tk.Frame(parent, bg=self.PANEL)
        title_row.pack(fill="x", padx=20, pady=(18, 8))
        tk.Label(title_row, text=self._t("inputs"), bg=self.PANEL, fg=self.TEXT,
                 font=("Segoe UI Semibold", 12)).pack(side="left")
        tk.Label(title_row, text=self._t("control"), bg=self.PANEL_ALT, fg=self.ACCENT_DARK, padx=10, pady=4,
                 font=("Consolas", 8, "bold")).pack(side="right")
        tk.Frame(parent, height=1, bg=self.BORDER).pack(fill="x", padx=20, pady=(0, 14))

        fields = tk.Frame(parent, bg=self.PANEL)
        fields.pack(fill="both", expand=True, padx=20)
        fields.grid_columnconfigure(0, weight=1)
        self._path_field(fields, 0, self._t("excel_required"), self.excel_file_var, self.choose_excel_file,
                         self._t("excel_hint"))

        sheet_row = tk.Frame(fields, bg=self.PANEL)
        sheet_row.grid(row=1, column=0, sticky="ew", pady=(0, 12))
        sheet_row.grid_columnconfigure(0, weight=1)
        tk.Label(sheet_row, text=self._t("sheet"), bg=self.PANEL, fg=self.MUTED,
                 font=("Segoe UI", 8, "bold")).grid(row=0, column=0, sticky="w", pady=(0, 5))
        self.sheet_combo = ttk.Combobox(sheet_row, textvariable=self.sheet_var, state="readonly", style="Light.TCombobox")
        self.sheet_combo.grid(row=1, column=0, sticky="ew")
        ttk.Button(sheet_row, text=self._t("reload"), style="Compact.TButton", command=self.load_sheets).grid(row=1, column=1, padx=(8, 0))
        if self.excel_file_var.get().strip():
            self.load_sheets(silent=True)

        self._path_field(fields, 2, self._t("word_optional"), self.word_file_var, self.choose_word_file,
                         self._t("word_hint"), clear_command=lambda: self.word_file_var.set(""))

        mode_row = tk.Frame(fields, bg=self.PANEL)
        mode_row.grid(row=3, column=0, sticky="ew", pady=(0, 12))
        tk.Label(mode_row, text=self._t("word_action"), bg=self.PANEL, fg=self.MUTED,
                 font=("Segoe UI", 8, "bold")).pack(side="left", padx=(0, 12))
        ttk.Radiobutton(mode_row, text=self._t("examine"), variable=self.word_mode_var, value="examine",
                        style="Mode.TRadiobutton").pack(side="left")
        ttk.Radiobutton(mode_row, text=self._t("fill"), variable=self.word_mode_var, value="fill",
                        style="Mode.TRadiobutton").pack(side="left", padx=(12, 0))

        self._path_field(fields, 4, self._t("pdf_optional"), self.pdf_folder_var, self.choose_pdf_folder,
                         self._t("pdf_hint"), clear_command=lambda: self.pdf_folder_var.set(""))

        issue_row = tk.Frame(fields, bg=self.PANEL)
        issue_row.grid(row=5, column=0, sticky="ew", pady=(0, 12))
        issue_row.grid_columnconfigure(0, weight=1)
        tk.Label(issue_row, text=self._t("expected_issue"), bg=self.PANEL, fg=self.MUTED,
                 font=("Segoe UI", 8, "bold")).grid(row=0, column=0, sticky="w", pady=(0, 5))
        self.ausgabe_entry = self._light_entry(issue_row, self.ausgabe_var)
        self.ausgabe_entry.grid(row=1, column=0, sticky="ew")
        tk.Label(issue_row, text=self._t("issue_hint"), bg=self.PANEL, fg=self.MUTED,
                 font=("Segoe UI", 8)).grid(row=2, column=0, sticky="w", pady=(5, 0))

        output_label = self._t("output_word") if self.word_mode_var.get() == "fill" else self._t("output_report")
        output_hint = self._t("output_word_hint") if self.word_mode_var.get() == "fill" else self._t("output_report_hint")
        self._path_field(fields, 6, output_label, self.output_file_var, self.choose_output_file, output_hint,
                         browse_text=self._t("save_as"))

        actions = tk.Frame(parent, bg=self.PANEL)
        actions.pack(fill="x", padx=20, pady=(0, 18))
        action_text = self._t("run_fill") if self.word_mode_var.get() == "fill" else self._t("run_check")
        self.run_button = ttk.Button(actions, text=action_text, style="Accent.TButton", command=self.run)
        self.run_button.pack(side="left")
        self.open_button = ttk.Button(actions, text=self._t("open_report"), style="Ghost.TButton",
                                      command=self.open_report, state="disabled")
        self.open_button.pack(side="left", padx=(9, 0))
        if self.word_mode_var.get() == "fill":
            self.open_button.pack_forget()

    def _build_results_panel(self, parent: tk.Frame):
        top = tk.Frame(parent, bg=self.PANEL)
        top.grid(row=0, column=0, sticky="ew", padx=16, pady=(16, 8))
        tk.Label(top, text=self._t("result"), bg=self.PANEL, fg=self.TEXT,
                 font=("Segoe UI Semibold", 11)).pack(side="left")
        self.activity_label = tk.Label(top, text=self._t("idle"), bg=self.PANEL_ALT, fg=self.MUTED,
                                       padx=9, pady=3, font=("Consolas", 8, "bold"))
        self.activity_label.pack(side="right")

        self.status_card = tk.Frame(parent, bg=self.PANEL_ALT, highlightthickness=1, highlightbackground=self.BORDER)
        self.status_card.grid(row=1, column=0, sticky="ew", padx=16, pady=(0, 10))
        self.status_card.grid_columnconfigure(0, weight=1)
        self.status_title_label = tk.Label(self.status_card, textvariable=self.result_title_var, bg=self.PANEL_ALT,
                                           fg=self.ACCENT_DARK, font=("Segoe UI Semibold", 15), anchor="w")
        self.status_title_label.grid(row=0, column=0, sticky="ew", padx=14, pady=(12, 2))
        self.status_detail_label = tk.Label(self.status_card, textvariable=self.result_detail_var, bg=self.PANEL_ALT,
                                            fg=self.TEXT, font=("Segoe UI", 9), anchor="w", justify="left", wraplength=520)
        self.status_detail_label.grid(row=1, column=0, sticky="ew", padx=14, pady=(0, 12))

        self.progress = ttk.Progressbar(parent, mode="determinate", value=0, maximum=100,
                                        style="Control.Horizontal.TProgressbar")
        self.progress.grid(row=2, column=0, sticky="ew", padx=16, pady=(0, 8))

        mode_strip = tk.Frame(parent, bg=self.PANEL_ALT, highlightthickness=1, highlightbackground=self.BORDER)
        mode_strip.grid(row=3, column=0, sticky="ew", padx=16, pady=(0, 10))
        tk.Label(mode_strip, text=self._t("mode"), bg=self.PANEL_ALT, fg=self.ACCENT,
                 font=("Consolas", 8, "bold")).pack(side="left", padx=(10, 8), pady=7)
        tk.Label(mode_strip, textvariable=self.mode_text_var, bg=self.PANEL_ALT, fg=self.TEXT,
                 font=("Segoe UI", 8), wraplength=470, justify="left").pack(side="left", fill="x", expand=True,
                                                                            padx=(0, 10), pady=7)

        table_frame = tk.Frame(parent, bg="#F8FBFC", highlightthickness=1, highlightbackground=self.BORDER)
        table_frame.grid(row=4, column=0, sticky="nsew", padx=16, pady=(0, 10))
        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)
        self.result_tree = ttk.Treeview(table_frame, columns=("id", "issue"), show="headings",
                                        style="Results.Treeview", selectmode="browse")
        self.result_tree.heading("id", text=self._t("id_heading"))
        self.result_tree.heading("issue", text=self._t("issue_heading"))
        self.result_tree.column("id", width=125, minwidth=90, stretch=False, anchor="w")
        self.result_tree.column("issue", width=410, minwidth=260, stretch=True, anchor="w")
        self.result_tree.grid(row=0, column=0, sticky="nsew")
        scroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.result_tree.yview)
        scroll.grid(row=0, column=1, sticky="ns")
        self.result_tree.configure(yscrollcommand=scroll.set)
        self.result_tree.tag_configure("error", background=self.ERROR_BG, foreground="#7D2630")
        self.result_tree.tag_configure("warning", background=self.WARNING_BG, foreground="#795000")
        self.result_tree.tag_configure("success", background=self.SUCCESS_BG, foreground=self.SUCCESS)
        self.result_tree.tag_configure("neutral", background="#F8FBFC", foreground=self.MUTED)

        footer = tk.Frame(parent, bg=self.PANEL)
        footer.grid(row=5, column=0, sticky="ew", padx=16, pady=(0, 16))
        footer.grid_columnconfigure(0, weight=1)
        tk.Label(footer, textvariable=self.current_step_var, bg=self.PANEL, fg=self.MUTED,
                 font=("Segoe UI", 8), anchor="w", justify="left", wraplength=400).grid(row=0, column=0, sticky="ew")
        self.correct_button = ttk.Button(footer, text=self._t("correct_copy"), style="Ghost.TButton",
                                         command=self.correct_word, state="disabled")
        self.correct_button.grid(row=0, column=1, padx=(10, 0))

    def _set_initial_result(self):
        self.result_title_var.set(self._t("await_title"))
        self.result_detail_var.set(self._t("await_detail"))
        self.current_step_var.set(self._t("idle"))
        if hasattr(self, "result_tree"):
            for item in self.result_tree.get_children():
                self.result_tree.delete(item)
            self.result_tree.insert("", "end", values=("—", self._t("no_result")), tags=("neutral",))

    def _light_entry(self, parent, variable: tk.StringVar) -> tk.Entry:
        return tk.Entry(parent, textvariable=variable, bg="#F8FBFC", fg=self.TEXT, insertbackground=self.ACCENT,
                        selectbackground=self.ACCENT_DARK, selectforeground="#FFFFFF", relief="flat",
                        highlightthickness=1, highlightbackground=self.BORDER, highlightcolor=self.ACCENT,
                        disabledbackground="#EEF3F5", disabledforeground=self.MUTED, font=("Segoe UI", 9))

    def _path_field(self, parent, row, label, variable, browse_command, file_hint, clear_command=None, browse_text=None):
        frame = tk.Frame(parent, bg=self.PANEL)
        frame.grid(row=row, column=0, sticky="ew", pady=(0, 12))
        frame.grid_columnconfigure(0, weight=1)
        tk.Label(frame, text=label, bg=self.PANEL, fg=self.MUTED, font=("Segoe UI", 8, "bold")).grid(row=0, column=0, sticky="w", pady=(0, 5))
        self._light_entry(frame, variable).grid(row=1, column=0, sticky="ew")
        ttk.Button(frame, text=browse_text or self._t("browse"), style="Compact.TButton", command=browse_command).grid(row=1, column=1, padx=(8, 0))
        if clear_command:
            ttk.Button(frame, text=self._t("clear"), style="Compact.TButton", command=clear_command).grid(row=1, column=2, padx=(6, 0))
        tk.Label(frame, text=file_hint, bg=self.PANEL, fg=self.MUTED, font=("Segoe UI", 8), wraplength=470,
                 justify="left").grid(row=2, column=0, columnspan=3, sticky="w", pady=(5, 0))

    def _change_language(self, _event=None):
        new_lang = "de" if self.language_var.get() == "Deutsch" else "en"
        if new_lang == self.lang:
            return
        self.lang = new_lang
        self._rebuild_ui()

    def _rebuild_ui(self):
        for child in list(self.winfo_children()):
            child.destroy()
        self._configure_styles()
        self._build_ui()
        self._refresh_mode_status()
        if self._last_summary:
            self._render_summary(self._last_summary)
        elif self._last_fill_result is not None:
            self._render_fill_result(self._last_fill_result)
        else:
            self._set_initial_result()
        if self._is_running:
            self.run_button.configure(state="disabled")
        elif self._last_run_succeeded and self.word_mode_var.get() == "examine":
            self.open_button.configure(state="normal")
            if self._last_summary.get("word_correction_needed", False):
                self.correct_button.configure(state="normal")

    def _on_word_mode_change(self):
        if not hasattr(self, "run_button"):
            return
        self._adjust_output_for_mode()
        self._last_run_succeeded = False
        self._last_summary = {}
        self._last_fill_result = None
        self._rebuild_ui()

    def _adjust_output_for_mode(self):
        current = Path(self.output_file_var.get().strip()) if self.output_file_var.get().strip() else None
        if self.word_mode_var.get() == "fill":
            word = self.word_file_var.get().strip()
            if word:
                suggested = Path(word).with_name(Path(word).stem + "_filled.docx")
            else:
                base = current.parent if current else Path.cwd()
                suggested = base / "filled_change_notice.docx"
            if not current or current.suffix.casefold() != ".docx":
                self.output_file_var.set(str(suggested))
        else:
            if not current or current.suffix.casefold() != ".xlsx":
                base = Path(self.excel_file_var.get().strip()).parent if self.excel_file_var.get().strip() else Path.cwd()
                self.output_file_var.set(str(base / "profile_docs_validation_report.xlsx"))

    def choose_pdf_folder(self):
        path = filedialog.askdirectory(title=self._t("choose_pdf_folder"))
        if path:
            self.pdf_folder_var.set(path)
            self._suggest_output(Path(path))

    def choose_excel_file(self):
        path = filedialog.askopenfilename(title=self._t("choose_excel_file"),
                                          filetypes=[("Excel", "*.xlsx *.xlsm"), ("All files", "*.*")])
        if path:
            self.excel_file_var.set(path)
            self.load_sheets()
            self._suggest_output(Path(path).parent)

    def choose_word_file(self):
        path = filedialog.askopenfilename(title=self._t("choose_word_file"),
                                          filetypes=[("Word", "*.docx"), ("All files", "*.*")])
        if path:
            self.word_file_var.set(path)
            if self.word_mode_var.get() == "fill":
                self.output_file_var.set(str(Path(path).with_name(Path(path).stem + "_filled.docx")))
            else:
                self._suggest_output(Path(path).parent)

    def choose_output_file(self):
        if self.word_mode_var.get() == "fill":
            initial = self.output_file_var.get().strip() or "filled_change_notice.docx"
            parent = Path(initial).parent
            path = filedialog.asksaveasfilename(title=self._t("choose_filled_word_path"), defaultextension=".docx",
                                                initialdir=str(parent) if parent.exists() else None,
                                                initialfile=Path(initial).name, filetypes=[("Word", "*.docx")])
        else:
            initial = self.output_file_var.get().strip() or "profile_docs_validation_report.xlsx"
            parent = Path(initial).parent
            path = filedialog.asksaveasfilename(title=self._t("choose_report_path"), defaultextension=".xlsx",
                                                initialdir=str(parent) if parent.exists() else None,
                                                initialfile=Path(initial).name, filetypes=[("Excel", "*.xlsx")])
        if path:
            self.output_file_var.set(path)

    def _suggest_output(self, base: Path):
        if self.output_file_var.get().strip():
            return
        folder = base if base.is_dir() else base.parent
        if self.word_mode_var.get() == "fill" and self.word_file_var.get().strip():
            word = Path(self.word_file_var.get().strip())
            self.output_file_var.set(str(folder / (word.stem + "_filled.docx")))
        elif self.word_mode_var.get() == "fill":
            self.output_file_var.set(str(folder / "filled_change_notice.docx"))
        else:
            self.output_file_var.set(str(folder / "profile_docs_validation_report.xlsx"))

    def load_sheets(self, silent: bool = False):
        excel = self.excel_file_var.get().strip()
        if not excel:
            return
        try:
            wb = load_workbook(excel, read_only=True, data_only=True)
            sheets = wb.sheetnames
            wb.close()
            self.sheet_combo["values"] = sheets
            if sheets and self.sheet_var.get() not in sheets:
                self.sheet_var.set(sheets[0])
            if not silent:
                self.current_step_var.set(self._t("sheets_loaded").format(sheets=", ".join(sheets)))
        except Exception as exc:
            if not silent:
                messagebox.showerror(self._t("sheet_load_error"), str(exc))

    def _refresh_mode_status(self):
        pdf = bool(self.pdf_folder_var.get().strip())
        word = bool(self.word_file_var.get().strip())
        fill = self.word_mode_var.get() == "fill"
        if fill:
            text = self._t("mode_fill_pdf") if pdf else self._t("mode_fill_no_pdf")
        elif pdf and word:
            text = self._t("mode_full")
        elif pdf:
            text = self._t("mode_pdf")
        elif word:
            text = self._t("mode_word")
        else:
            text = self._t("select_sources")
        self.mode_text_var.set(text)
        if hasattr(self, "ausgabe_entry"):
            self.ausgabe_entry.configure(state="normal" if pdf and not fill else "disabled")

    def _collect_args(self):
        pdf_folder = self.pdf_folder_var.get().strip() or None
        excel = self.excel_file_var.get().strip()
        word_file = self.word_file_var.get().strip() or None
        output = self.output_file_var.get().strip()
        ausgabe = self.ausgabe_var.get().strip() or None
        if not excel:
            raise ValueError(self._t("choose_excel"))
        if self.word_mode_var.get() == "fill":
            if not word_file:
                raise ValueError(self._t("choose_word_fill"))
        elif not pdf_folder and not word_file:
            raise ValueError(self._t("choose_source"))
        if self.word_mode_var.get() == "examine" and pdf_folder and not ausgabe:
            raise ValueError(self._t("enter_issue"))
        if not output:
            raise ValueError(self._t("choose_output"))
        return pdf_folder, excel, word_file, output, ausgabe

    def run(self):
        try:
            args = self._collect_args()
        except Exception as exc:
            messagebox.showerror(self._t("missing_input"), str(exc))
            return
        pdf_folder, _excel, word_file, _output, _ausgabe = args
        fill = self.word_mode_var.get() == "fill"
        if word_file and not pdf_folder:
            question = self._t("page_unavailable_fill") if fill else self._t("page_unavailable_check")
            if not messagebox.askyesno(self._t("page_unavailable_title"), question, icon="warning"):
                return
        self._begin_run(fill=fill)
        if fill:
            threading.Thread(target=self._fill_thread, args=args, daemon=True).start()
        else:
            threading.Thread(target=self._run_thread, args=args, daemon=True).start()

    def _begin_run(self, fill: bool):
        self._is_running = True
        self._last_run_succeeded = False
        self._last_summary = {}
        self._last_fill_result = None
        self.open_button.configure(state="disabled")
        self.correct_button.configure(state="disabled")
        self.run_button.configure(state="disabled")
        self.activity_label.configure(text=self._t("running"), fg=self.ACCENT)
        self.result_title_var.set(self._t("validation_progress") if not fill else self._t("run_fill"))
        self.result_detail_var.set(self._t("validation_detail"))
        self.status_card.configure(bg=self.PANEL_ALT, highlightbackground=self.BORDER)
        self.status_title_label.configure(bg=self.PANEL_ALT, fg=self.ACCENT_DARK)
        self.status_detail_label.configure(bg=self.PANEL_ALT)
        self.progress.configure(mode="indeterminate")
        self.progress.start(12)
        for item in self.result_tree.get_children():
            self.result_tree.delete(item)
        self.result_tree.insert("", "end", values=("…", self._t("checking")), tags=("neutral",))
        self.current_step_var.set(self._t("start_validation"))

    def _translate_progress(self, message: str) -> str:
        if self.lang == "en":
            return message
        replacements = [
            ("Loading Excel...", "Excel wird geladen…"),
            ("Loading Excel rows for Word filling...", "Excel-Zeilen zum Befüllen werden geladen…"),
            ("Recognizing the Freigabe-/Änderungsmitteilung Word template...", "Word-Änderungsmitteilung wird erkannt…"),
            ("Comparing Word rows with Excel...", "Word-Zeilen werden mit Excel verglichen…"),
            ("No PDF folder selected.", "Kein PDF-Ordner ausgewählt."),
            ("Processing PDF", "PDF wird verarbeitet"),
            ("Reading PDF page count", "PDF-Seitenzahl wird gelesen"),
            ("Filling Word row", "Word-Zeile wird befüllt"),
            ("Done. Report saved to:", "Fertig. Bericht gespeichert unter:"),
            ("Filled Word copy saved:", "Befüllte Word-Kopie gespeichert:"),
            ("Corrected Word copy saved:", "Korrigierte Word-Kopie gespeichert:"),
            ("Found", "Gefunden:"),
            ("Loaded", "Geladen:"),
        ]
        translated = message
        for old, new in replacements:
            translated = translated.replace(old, new)
        return translated

    def _run_thread(self, pdf_folder, excel, word_file, output, ausgabe):
        summary_holder: dict = {}
        try:
            def progress(message: str):
                self.after(0, lambda m=self._translate_progress(message): self.current_step_var.set(m))
            run_validation(pdf_folder=pdf_folder, excel_path=excel, output_path=output,
                           expected_ausgabe=ausgabe, sheet_name=self.sheet_var.get().strip() or None,
                           page_number=2, word_docx_path=word_file, progress_callback=progress,
                           summary_callback=lambda summary: summary_holder.update(summary))
            self.after(0, lambda: self._finish_success(output, summary_holder))
        except Exception as exc:
            self.after(0, lambda: self._finish_error(str(exc)))

    def _fill_thread(self, pdf_folder, excel, word_file, output, _ausgabe):
        try:
            def progress(message: str):
                self.after(0, lambda m=self._translate_progress(message): self.current_step_var.set(m))
            result = fill_word_copy_from_excel(word_docx_path=word_file, excel_path=excel, output_path=output,
                                               sheet_name=self.sheet_var.get().strip() or None,
                                               pdf_folder=pdf_folder, progress_callback=progress)
            self.after(0, lambda: self._finish_fill(result))
        except Exception as exc:
            self.after(0, lambda: self._finish_fill_error(str(exc)))

    def _finish_success(self, output: str, summary: dict):
        self._is_running = False
        self.progress.stop()
        self.progress.configure(mode="determinate", value=100)
        self.run_button.configure(state="normal")
        self.open_button.configure(state="normal")
        self.activity_label.configure(text=self._t("complete"), fg=self.SUCCESS)
        self.current_step_var.set(self._t("report_saved").format(path=output))
        self._last_run_succeeded = True
        self._last_summary = summary
        self._render_summary(summary)
        should_enable = bool(self.word_file_var.get().strip()) and bool(summary.get("word_correction_needed", False))
        self.correct_button.configure(state="normal" if should_enable else "disabled")

    def _localize_issue(self, text: str) -> str:
        if self.lang == "en":
            return text
        replacements = {
            "User input": "Benutzereingabe",
            "<empty>": "<leer>",
            "unavailable": "nicht verfügbar",
            "Word row": "Word-Zeile",
            "PDF folder": "PDF-Ordner",
            "Word table": "Word-Tabelle",
            "could not be verified because no PDF folder was selected": "konnte ohne ausgewählten PDF-Ordner nicht geprüft werden",
            "PDF not found for Dokumentnummer": "PDF nicht gefunden für Dokumentnummer",
            "Word entry not found for Dokumentnummer": "Word-Eintrag nicht gefunden für Dokumentnummer",
            "No reliable Excel match": "Keine zuverlässige Excel-Zuordnung",
            "Document not found in Excel": "Dokument nicht in Excel gefunden",
            "Ambiguous Excel match": "Mehrdeutige Excel-Zuordnung",
        }
        for old, new in replacements.items():
            text = text.replace(old, new)
        return text

    def _render_summary(self, summary: dict):
        for item in self.result_tree.get_children():
            self.result_tree.delete(item)
        issues = summary.get("issues", [])
        checked = int(summary.get("checked_entries", 0))
        if summary.get("all_correct", False):
            self.result_title_var.set(self._t("all_correct"))
            entry = self._t("entry_singular") if checked == 1 else self._t("entry_plural")
            self.result_detail_var.set(self._t("all_correct_detail").format(count=checked, entry=entry))
            self.status_card.configure(bg=self.SUCCESS_BG, highlightbackground="#A9D8C1")
            self.status_title_label.configure(bg=self.SUCCESS_BG, fg=self.SUCCESS)
            self.status_detail_label.configure(bg=self.SUCCESS_BG)
            self.result_tree.insert("", "end", values=("✓", self._t("all_values_correct")), tags=("success",))
            if hasattr(self, "correct_button"):
                self.correct_button.configure(state="disabled")
            return
        errors = int(summary.get("error_count", 0))
        warnings = int(summary.get("warning_count", 0))
        item_word = self._t("item_singular") if len(issues) == 1 else self._t("item_plural")
        self.result_title_var.set(self._t("items_attention").format(count=len(issues), items=item_word))
        error_word = self._t("error_singular") if errors == 1 else self._t("error_plural")
        warnings_text = ""
        if warnings:
            warning_word = self._t("warning_singular") if warnings == 1 else self._t("warning_plural")
            warnings_text = (f" und {warnings} {warning_word}" if self.lang == "de" else f" and {warnings} {warning_word}")
        self.result_detail_var.set(self._t("errors_warnings").format(errors=errors, error_word=error_word,
                                                                      warnings_text=warnings_text))
        self.status_card.configure(bg=self.ERROR_BG if errors else self.WARNING_BG,
                                   highlightbackground="#E4B8BE" if errors else "#E5CB88")
        self.status_title_label.configure(bg=self.ERROR_BG if errors else self.WARNING_BG,
                                          fg=self.ERROR if errors else self.WARNING)
        self.status_detail_label.configure(bg=self.ERROR_BG if errors else self.WARNING_BG)
        for issue in issues:
            severity = issue.get("severity", "ERROR")
            dok_id = issue.get("dok_id") or "—"
            text = self._localize_issue(issue.get("issue", ""))
            source = self._localize_issue(issue.get("source", ""))
            if source:
                text = f"{text}  [{source}]"
            self.result_tree.insert("", "end", values=(dok_id, text),
                                    tags=("warning" if severity == "WARNING" else "error",))

    def _finish_error(self, error_text: str):
        self._is_running = False
        self.progress.stop()
        self.progress.configure(mode="determinate", value=0)
        self.run_button.configure(state="normal")
        self.activity_label.configure(text=self._t("error"), fg=self.ERROR)
        self.result_title_var.set(self._t("validation_failed"))
        self.result_detail_var.set(error_text)
        self.status_card.configure(bg=self.ERROR_BG, highlightbackground="#E4B8BE")
        self.status_title_label.configure(bg=self.ERROR_BG, fg=self.ERROR)
        self.status_detail_label.configure(bg=self.ERROR_BG)
        self.current_step_var.set(self._t("validation_did_not_finish"))
        for item in self.result_tree.get_children():
            self.result_tree.delete(item)
        self.result_tree.insert("", "end", values=("—", error_text), tags=("error",))
        messagebox.showerror(self._t("validation_failed"), error_text)

    def _finish_fill(self, result):
        self._is_running = False
        self.progress.stop()
        self.progress.configure(mode="determinate", value=100)
        self.run_button.configure(state="normal")
        self.activity_label.configure(text=self._t("complete"), fg=self.SUCCESS)
        self.current_step_var.set(self._t("filled_saved").format(path=result.output_path))
        self._last_fill_result = result
        self._render_fill_result(result)

    def _render_fill_result(self, result):
        for item in self.result_tree.get_children():
            self.result_tree.delete(item)
        self.result_title_var.set(self._t("filled_title"))
        self.result_detail_var.set(self._t("filled_detail").format(rows=result.filled_rows))
        if result.warnings:
            self.status_card.configure(bg=self.WARNING_BG, highlightbackground="#E5CB88")
            self.status_title_label.configure(bg=self.WARNING_BG, fg=self.WARNING)
            self.status_detail_label.configure(bg=self.WARNING_BG)
            for warning in result.warnings:
                self.result_tree.insert("", "end", values=("—", self._t("filled_warning").format(warning=self._localize_issue(warning))),
                                        tags=("warning",))
        else:
            self.status_card.configure(bg=self.SUCCESS_BG, highlightbackground="#A9D8C1")
            self.status_title_label.configure(bg=self.SUCCESS_BG, fg=self.SUCCESS)
            self.status_detail_label.configure(bg=self.SUCCESS_BG)
            self.result_tree.insert("", "end", values=("✓", self._t("filled_success")), tags=("success",))
        self.correct_button.configure(state="disabled")

    def _finish_fill_error(self, error_text: str):
        self._is_running = False
        self.progress.stop()
        self.progress.configure(mode="determinate", value=0)
        self.run_button.configure(state="normal")
        self.activity_label.configure(text=self._t("error"), fg=self.ERROR)
        self.result_title_var.set(self._t("fill_failed"))
        self.result_detail_var.set(error_text)
        self.current_step_var.set(self._t("fill_did_not_finish"))
        self.status_card.configure(bg=self.ERROR_BG, highlightbackground="#E4B8BE")
        self.status_title_label.configure(bg=self.ERROR_BG, fg=self.ERROR)
        self.status_detail_label.configure(bg=self.ERROR_BG)
        for item in self.result_tree.get_children():
            self.result_tree.delete(item)
        self.result_tree.insert("", "end", values=("—", error_text), tags=("error",))
        messagebox.showerror(self._t("fill_failed"), error_text)

    def correct_word(self):
        if (not self._last_run_succeeded or not self.word_file_var.get().strip()
                or not self._last_summary.get("word_correction_needed", False)):
            messagebox.showerror(self._t("no_word_result"), self._t("no_word_result_msg"))
            return
        source = Path(self.word_file_var.get().strip())
        suggested = source.with_name(source.stem + "_corrected.docx")
        output = filedialog.asksaveasfilename(title=self._t("save_corrected"), defaultextension=".docx",
                                              initialdir=str(suggested.parent), initialfile=suggested.name,
                                              filetypes=[("Word", "*.docx")])
        if not output:
            return
        if not messagebox.askyesno(self._t("create_corrected_title"), self._t("create_corrected_question"), icon="question"):
            return
        self.correct_button.configure(state="disabled")
        self.run_button.configure(state="disabled")
        self.activity_label.configure(text=self._t("correcting"), fg=self.ACCENT)
        self.current_step_var.set(self._t("creating_corrected"))
        threading.Thread(target=self._correct_word_thread, args=(output,), daemon=True).start()

    def _correct_word_thread(self, output: str):
        try:
            def progress(message: str):
                self.after(0, lambda m=self._translate_progress(message): self.current_step_var.set(m))
            result = create_corrected_word_copy(word_docx_path=self.word_file_var.get().strip(),
                                                excel_path=self.excel_file_var.get().strip(), output_path=output,
                                                sheet_name=self.sheet_var.get().strip() or None,
                                                pdf_folder=self.pdf_folder_var.get().strip() or None,
                                                progress_callback=progress)
            self.after(0, lambda: self._finish_correction(result))
        except Exception as exc:
            self.after(0, lambda: self._finish_correction_error(str(exc)))

    def _finish_correction(self, result):
        self.run_button.configure(state="normal")
        self.correct_button.configure(state="disabled")
        self.activity_label.configure(text=self._t("complete"), fg=self.SUCCESS)
        self.current_step_var.set(self._t("corrected_saved").format(path=result.output_path))
        warning_text = ""
        if result.warnings:
            warning_text = "\n\n" + self._t("warnings") + ":\n- " + "\n- ".join(result.warnings)
        messagebox.showinfo(self._t("corrected_created"),
                            self._t("corrected_info").format(path=result.output_path, cells=result.changed_cells,
                                                              rows=result.corrected_rows, skipped=result.skipped_rows,
                                                              warnings=warning_text))

    def _finish_correction_error(self, error_text: str):
        self.run_button.configure(state="normal")
        self.correct_button.configure(state="normal")
        self.activity_label.configure(text=self._t("error"), fg=self.ERROR)
        self.current_step_var.set(self._t("could_not_correct"))
        messagebox.showerror(self._t("correction_failed"), error_text)

    def open_report(self):
        path = self.output_file_var.get().strip()
        if not path or not Path(path).exists():
            messagebox.showerror(self._t("report_not_found"), self._t("report_not_found_msg"))
            return
        try:
            if os.name == "nt":
                os.startfile(path)  # type: ignore[attr-defined]
            else:
                webbrowser.open(Path(path).resolve().as_uri())
        except Exception as exc:
            messagebox.showerror(self._t("could_not_open_report"), str(exc))


def main():
    app = App()
    app.mainloop()


if __name__ == "__main__":
    main()
