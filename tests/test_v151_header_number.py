from __future__ import annotations

from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.shared import Inches

from doc_table_checker.core import (
    create_corrected_word_copy,
    extract_change_notice_word_records,
)
from test_synthetic import make_dataset


def make_header_number_change_notice(
    path: Path,
    header_number: str,
    body_number: str = "WRONG-BODY-NUMBER",
    wrong_values: bool = False,
):
    doc = Document()

    # Exact user-provided address:
    # section[0].header.table[0].row[0].cell[1].paragraph[0].word[0]
    header_table = doc.sections[0].header.add_table(rows=1, cols=2, width=Inches(6.5))
    header_table.cell(0, 0).text = "Header label"
    header_table.cell(0, 1).text = header_number

    # The standard body table still contains a misleading Number-like value to
    # prove the exact header address takes precedence.
    table = doc.add_table(rows=5, cols=13)
    table.cell(0, 0).merge(table.cell(0, 11)).text = (
        "Freigabe-/Änderungsmitteilung / Engineering Change Notice"
    )
    table.cell(0, 12).text = body_number
    table.cell(1, 0).merge(table.cell(1, 12)).text = "Benennung: / Description:"

    table.cell(2, 0).merge(table.cell(3, 0)).text = "Artikel-Nr. / part no."
    table.cell(2, 1).merge(table.cell(2, 2)).text = "o. Zeichnung / w/o drawing"
    table.cell(2, 3).merge(table.cell(3, 3)).text = "Dokument-Nr. / document no."
    table.cell(2, 4).merge(table.cell(2, 8)).text = (
        "Zeichnung / Dokument / Stückliste / drawing / document / parts list"
    )
    table.cell(2, 9).merge(table.cell(3, 9)).text = "Blatt / sheets"
    table.cell(2, 10).merge(table.cell(3, 10)).text = "Verfügung / disposal"
    table.cell(2, 11).merge(table.cell(3, 11)).text = (
        "Zulassungs-/ relevant / approval relevant"
    )
    table.cell(2, 12).merge(table.cell(3, 12)).text = "Beschreibung / description:"
    table.cell(3, 1).text = "Rev alt/old"
    table.cell(3, 2).text = "Rev neu/new"
    table.cell(3, 4).text = "Format"
    table.cell(3, 5).text = "Rev alt/old"
    table.cell(3, 6).text = "Vers alt/old"
    table.cell(3, 7).text = "Rev neu/new"
    table.cell(3, 8).text = "Vers neu/new"

    table.cell(4, 0).text = "WRONG-ARTICLE" if wrong_values else "ART-123"
    table.cell(4, 3).text = "EE******"
    table.cell(4, 7).text = "99" if wrong_values else "01"
    table.cell(4, 8).text = "88" if wrong_values else "02"
    table.cell(4, 9).text = "9" if wrong_values else "2"
    doc.save(path)


def run_test():
    base = ROOT / "test_output_v151"
    pdf_dir, excel = make_dataset(base)
    word = base / "exact_header_number.docx"
    make_header_number_change_notice(word, "FR-1", body_number="WRONG-BODY-NUMBER")

    records, info = extract_change_notice_word_records(word)
    assert info.change_number_raw == "FR-1", info.change_number_raw
    assert info.change_number_source == (
        "section[0].header.table[0].row[0].cell[1].paragraph[0].word[0]"
    )
    assert info.change_number_section_zero_based == 0
    assert info.change_number_header_table_zero_based == 0
    assert info.change_number_cell_row_zero_based == 0
    assert info.change_number_cell_col_zero_based == 1
    assert records[0].values_raw["freigabe"] == "FR-1"

    wrong_word = base / "wrong_exact_header_number.docx"
    make_header_number_change_notice(
        wrong_word,
        "WRONG-HEADER-NUMBER",
        body_number="FR-1",
        wrong_values=True,
    )
    corrected = base / "wrong_exact_header_number_corrected.docx"
    correction = create_corrected_word_copy(
        wrong_word, excel, corrected, sheet_name="Master", pdf_folder=pdf_dir,
    )
    assert correction.changed_cells >= 5, correction

    corrected_doc = Document(corrected)
    corrected_header_cell = corrected_doc.sections[0].header.tables[0].rows[0].cells[1]
    assert corrected_header_cell.paragraphs[0].text.split()[0] == "FR-1"
    corrected_records, corrected_info = extract_change_notice_word_records(corrected)
    assert corrected_info.change_number_raw == "FR-1"
    assert corrected_records[0].values_raw["freigabe"] == "FR-1"
    print("v1.5.1 exact header Number test passed")


if __name__ == "__main__":
    run_test()
