from __future__ import annotations

from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from openpyxl import load_workbook

from doc_table_checker.core import run_validation
from test_synthetic import make_dataset
from test_change_notice_auto import make_change_notice


def run_word_only_test():
    base = ROOT / "test_output_word_only"
    _pdf_dir, excel = make_dataset(base)
    word = base / "change_notice_word_only.docx"
    make_change_notice(word)
    out = base / "word_only_report.xlsx"

    results = run_validation(
        pdf_folder=None,
        excel_path=excel,
        output_path=out,
        expected_ausgabe=None,
        sheet_name="Master",
        word_docx_path=word,
    )
    assert results == []
    assert out.exists()

    wb = load_workbook(out, read_only=True, data_only=True)
    assert "Word page count" in wb.sheetnames
    page_results = [row[8] for row in wb["Word page count"].iter_rows(min_row=2, values_only=True)]
    assert page_results and set(page_results) == {"UNAVAILABLE_NO_PDFS"}
    run_info = {row[0]: row[1] for row in wb["Run info"].iter_rows(min_row=2, values_only=True)}
    assert run_info["PDF folder provided"] == "No"
    assert "Unavailable" in run_info["Page-count verification"]
    wb.close()
    print("Word-only validation test passed.")


if __name__ == "__main__":
    run_word_only_test()
