from __future__ import annotations

from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from openpyxl import load_workbook

from doc_table_checker.core import write_word_inspection_files, load_word_mapping, extract_word_records, run_validation
from test_synthetic import make_dataset


def run_word_synthetic_test():
    base = ROOT / "test_output"
    pdf_dir, excel = make_dataset(base)

    word = base / "word_list.docx"
    doc = Document()
    doc.add_paragraph("Synthetic Word document list")
    table = doc.add_table(rows=1, cols=7)
    headers = ["Sprache", "Dok.-Nr.", "DOK-ID", "Freigabe-/ Änd.-Nr.", "Artikel-Nr.", "Rev.", "Vers."]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    rows = [
        ["de", "EE******", "DOC-001", "FR-1", "ART-123", "01", "02"],
        ["en", "DN-200", "DOC-002", "FR-2", "ART-222", "1", "1"],
        ["fr", "DN-400", "DOC-WRONG", "FR-5", "ART-555", "03", "02"],
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    doc.save(word)

    structure = base / "word_table_structure.json"
    mapping = base / "word_mapping_template.json"
    write_word_inspection_files(word, structure, mapping)
    assert structure.exists()
    assert mapping.exists()

    word_mapping = load_word_mapping(mapping)
    word_records = extract_word_records(word, word_mapping)
    assert len(word_records) == 3
    assert word_records[0].values_norm["dok_id"] == "DOC-001"

    out = base / "validation_report_with_word.xlsx"
    run_validation(
        pdf_dir,
        excel,
        out,
        expected_ausgabe="07-2026",
        sheet_name="Master",
        page_number=2,
        word_docx_path=word,
        word_mapping_path=mapping,
    )
    assert out.exists()
    wb = load_workbook(out, read_only=True, data_only=True)
    assert wb.sheetnames == ["Overview", "Issues"]
    issues = list(wb["Issues"].iter_rows(min_row=5, values_only=True))
    assert any(row[5] == "Word matching" and "DOK-ID" in str(row[11]) for row in issues)
    wb.close()
    print("Word synthetic test passed.")
    print("Sample report:", out)


if __name__ == "__main__":
    run_word_synthetic_test()
