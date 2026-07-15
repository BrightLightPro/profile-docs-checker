from __future__ import annotations

from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from openpyxl import load_workbook

from doc_table_checker.core import extract_change_notice_word_records, run_validation
from test_synthetic import make_dataset


def make_change_notice(path: Path):
    doc = Document()
    table = doc.add_table(rows=6, cols=13)

    table.cell(0, 0).merge(table.cell(0, 11)).text = (
        "Freigabe-/Änderungsmitteilung / Engineering Change Notice"
    )
    table.cell(0, 12).text = "FR-1"
    table.cell(1, 0).merge(table.cell(1, 12)).text = "Benennung: / Description:"

    table.cell(2, 0).merge(table.cell(3, 0)).text = "Artikel-Nr. / part no."
    table.cell(2, 1).merge(table.cell(2, 2)).text = "o. Zeichnung / w/o drawing"
    table.cell(2, 3).merge(table.cell(3, 3)).text = "Dokument-Nr. / document no."
    table.cell(2, 4).merge(table.cell(2, 8)).text = (
        "Zeichnung / Dokument / Stückliste / drawing / document / parts list"
    )
    table.cell(2, 9).merge(table.cell(3, 9)).text = "Blatt / sheets"
    table.cell(2, 10).merge(table.cell(3, 10)).text = "Verfügung / disposal"
    table.cell(2, 11).merge(table.cell(3, 11)).text = (
        "Zulassungs-/ relevant / approval relevant"
    )
    table.cell(2, 12).merge(table.cell(3, 12)).text = "Beschreibung / description:"

    table.cell(3, 1).text = "Rev alt/old"
    table.cell(3, 2).text = "Rev neu/new"
    table.cell(3, 4).text = "Format"
    table.cell(3, 5).text = "Rev alt/old"
    table.cell(3, 6).text = "Vers alt/old"
    table.cell(3, 7).text = "Rev neu/new"
    table.cell(3, 8).text = "Vers neu/new"

    rows = [
        ("ART-123", "EE******", "01", "02", "3"),  # page mismatch; PDF has 2 pages
        ("ART-222", "DN-200", "1", "1", "2"),      # top-right number differs from Excel FR-2
    ]
    for row_idx, (article, doc_no, rev_new, vers_new, sheets) in enumerate(rows, start=4):
        table.cell(row_idx, 0).text = article
        table.cell(row_idx, 3).text = doc_no
        table.cell(row_idx, 7).text = rev_new
        table.cell(row_idx, 8).text = vers_new
        table.cell(row_idx, 9).text = sheets

    doc.save(path)


def run_change_notice_test():
    base = ROOT / "test_output"
    pdf_dir, excel = make_dataset(base)
    word = base / "change_notice.docx"
    make_change_notice(word)

    records, info = extract_change_notice_word_records(word)
    assert len(records) == 2
    assert info.table_index_zero_based == 0
    assert info.data_start_row_zero_based == 4
    assert info.columns_by_index_zero_based == {
        "artikelnummer": 0,
        "dokumentnummer": 3,
        "revision": 7,
        "version": 8,
        "seiten": 9,
    }
    assert info.change_number_raw == "FR-1"
    assert records[0].values_norm["revision"] == "01"
    assert records[0].values_norm["version"] == "02"
    assert "zulassungs" not in records[0].values_raw

    out = base / "validation_report_auto_word.xlsx"
    run_validation(
        pdf_dir,
        excel,
        out,
        expected_ausgabe="07-2026",
        sheet_name="Master",
        page_number=2,
        word_docx_path=word,
    )
    assert out.exists()
    wb = load_workbook(out, read_only=True, data_only=True)
    assert wb.sheetnames == ["Overview", "Issues"]
    issues = list(wb["Issues"].iter_rows(min_row=5, values_only=True))
    assert any(row[5] == "Word vs PDF" and row[6] == "Blatt / sheets" for row in issues)
    assert any(row[5] == "Word vs Excel" and row[6] == "Freigabe-/Änd.-Nr." for row in issues)
    flattened = [value for row in issues for value in row if value is not None]
    assert not any("approval relevant" in str(value).casefold() for value in flattened)
    wb.close()
    print("Automatic change-notice Word test passed.")
    print("Sample report:", out)


if __name__ == "__main__":
    run_change_notice_test()
