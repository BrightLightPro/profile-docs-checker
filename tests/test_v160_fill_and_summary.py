from __future__ import annotations

from pathlib import Path
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook

from doc_table_checker.core import (
    fill_word_copy_from_excel,
    extract_change_notice_word_records,
)


def make_excel(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append(["Sprache", "Dok.-Nr.", "DOK-ID", "Freigabe-/ Änd.-Nr.", "Artikel-Nr.", "Rev.", "Vers."])
    ws.append(["de", "EE100", "ID-100", "ECN-77", "A-100", "01", "02"])
    ws.append(["en", "EE200", "ID-200", "ECN-77", "A-200", "03", "04"])
    ws.append(["jp", "EE300", "ID-300", "ECN-77", "A-300", "A", "B"])
    wb.save(path)


def make_blank_template(path: Path):
    doc = Document()
    header_table = doc.sections[0].header.add_table(rows=1, cols=2, width=Inches(6.5))
    header_table.cell(0, 0).text = "Header label"
    header_table.cell(0, 1).text = "(Number)"

    table = doc.add_table(rows=5, cols=13)
    table.cell(0, 0).merge(table.cell(0, 11)).text = "Freigabe-/Änderungsmitteilung / Engineering Change Notice"
    table.cell(0, 12).text = "(Number)"
    table.cell(1, 0).merge(table.cell(1, 12)).text = "Benennung: / Description:"
    table.cell(2, 0).merge(table.cell(3, 0)).text = "Artikel-Nr. / part no."
    table.cell(2, 1).merge(table.cell(2, 2)).text = "o. Zeichnung / w/o drawing"
    table.cell(2, 3).merge(table.cell(3, 3)).text = "Dokument-Nr. / document no."
    table.cell(2, 4).merge(table.cell(2, 8)).text = "Zeichnung / Dokument / Stückliste / drawing / document / parts list"
    table.cell(2, 9).merge(table.cell(3, 9)).text = "Blatt / sheets"
    table.cell(2, 10).merge(table.cell(3, 10)).text = "Verfügung / disposal"
    table.cell(2, 11).merge(table.cell(3, 11)).text = "Zulassungs-/ relevant / approval relevant"
    table.cell(2, 12).merge(table.cell(3, 12)).text = "Beschreibung / description:"
    table.cell(3, 1).text = "Rev alt/old"
    table.cell(3, 2).text = "Rev neu/new"
    table.cell(3, 4).text = "Format"
    table.cell(3, 5).text = "Rev alt/old"
    table.cell(3, 6).text = "Vers alt/old"
    table.cell(3, 7).text = "Rev neu/new"
    table.cell(3, 8).text = "Vers neu/new"
    doc.save(path)


def run_test():
    base = ROOT / "test_output_v160"
    if base.exists():
        shutil.rmtree(base)
    base.mkdir()
    excel = base / "master.xlsx"
    template = base / "blank.docx"
    output = base / "filled.docx"
    make_excel(excel)
    make_blank_template(template)

    result = fill_word_copy_from_excel(template, excel, output, sheet_name="Master")
    assert output.exists()
    assert result.filled_rows == 3
    assert result.added_rows == 2
    assert any("No PDF folder" in warning for warning in result.warnings)

    records, info = extract_change_notice_word_records(output)
    assert info.change_number_raw == "ECN-77"
    assert len(records) == 3
    assert records[0].values_raw["artikelnummer"] == "A-100"
    assert records[0].values_raw["dokumentnummer"] == "EE100"
    assert records[0].values_raw["revision"] == "01"
    assert records[0].values_raw["version"] == "02"
    assert records[0].values_raw["seiten"] == ""
    assert records[2].values_raw["dokumentnummer"] == "EE300"
    print("v1.6 fill test passed")


if __name__ == "__main__":
    run_test()
