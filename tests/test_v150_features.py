from __future__ import annotations

from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document

from doc_table_checker.core import (
    create_corrected_word_copy,
    extract_change_notice_word_records,
    run_validation,
)
from test_synthetic import make_dataset


def make_single_change_notice(path: Path, number: str, distractor: str = "", wrong_values: bool = False):
    doc = Document()
    table = doc.add_table(rows=5, cols=13)
    table.cell(0, 0).merge(table.cell(0, 11)).text = "Freigabe-/Änderungsmitteilung / Engineering Change Notice"
    table.cell(0, 12).text = number
    table.cell(1, 0).merge(table.cell(1, 12)).text = "Benennung: / Description: " + distractor

    table.cell(2, 0).merge(table.cell(3, 0)).text = "Artikel-Nr. / part no."
    table.cell(2, 1).merge(table.cell(2, 2)).text = "o. Zeichnung / w/o drawing"
    table.cell(2, 3).merge(table.cell(3, 3)).text = "Dokument-Nr. / document no."
    table.cell(2, 4).merge(table.cell(2, 8)).text = "Zeichnung / Dokument / Stückliste / drawing / document / parts list"
    table.cell(2, 9).merge(table.cell(3, 9)).text = "Blatt / sheets"
    table.cell(2, 10).merge(table.cell(3, 10)).text = "Verfügung / disposal"
    table.cell(2, 11).merge(table.cell(3, 11)).text = "Zulassungs-/ relevant / approval relevant"
    table.cell(2, 12).merge(table.cell(3, 12)).text = "Beschreibung / description:"
    table.cell(3, 1).text = "Rev alt/old"
    table.cell(3, 2).text = "Rev neu/new"
    table.cell(3, 4).text = "Format"
    table.cell(3, 5).text = "Rev alt/old"
    table.cell(3, 6).text = "Vers alt/old"
    table.cell(3, 7).text = "Rev neu/new"
    table.cell(3, 8).text = "Vers neu/new"

    table.cell(4, 0).text = "WRONG-ARTICLE" if wrong_values else "ART-123"
    table.cell(4, 3).text = "EE******"
    table.cell(4, 7).text = "99" if wrong_values else "01"
    table.cell(4, 8).text = "88" if wrong_values else "02"
    table.cell(4, 9).text = "9" if wrong_values else "2"
    doc.save(path)


def run_test():
    base = ROOT / "test_output_v150"
    pdf_dir, excel = make_dataset(base)

    # The extractor must use the dedicated top-right cell, not another code in the header.
    word = base / "number_location.docx"
    make_single_change_notice(word, "(Number) FR-1", distractor="DISTRACTOR-999")
    records, info = extract_change_notice_word_records(word)
    assert info.change_number_raw == "FR-1", info.change_number_raw
    assert info.change_number_cell_row_zero_based == 0
    assert info.change_number_cell_col_zero_based == 12
    assert records[0].values_raw["freigabe"] == "FR-1"

    # The right-panel callback must receive concise actionable items.
    wrong_word = base / "wrong_change_notice.docx"
    make_single_change_notice(wrong_word, "WRONG-NUMBER", wrong_values=True)
    summary = {}
    report = base / "v150_report.xlsx"
    run_validation(
        pdf_dir, excel, report, expected_ausgabe="07.2026", sheet_name="Master",
        word_docx_path=wrong_word, summary_callback=lambda value: summary.update(value),
    )
    assert summary["issue_count"] > 0
    assert any("Freigabe-/Änd.-Nr." in item["issue"] for item in summary["issues"])
    assert any(item["dok_id"] == "DOC-001" for item in summary["issues"])

    corrected = base / "wrong_change_notice_corrected.docx"
    correction = create_corrected_word_copy(
        wrong_word, excel, corrected, sheet_name="Master", pdf_folder=pdf_dir,
    )
    assert correction.changed_cells >= 5, correction
    assert corrected.exists()

    corrected_records, corrected_info = extract_change_notice_word_records(corrected)
    row = corrected_records[0].values_raw
    assert corrected_info.change_number_raw == "FR-1"
    assert row["artikelnummer"] == "ART-123"
    assert row["dokumentnummer"] == "EE******"
    assert row["revision"] == "01"
    assert row["version"] == "02"
    assert row["seiten"] == "2"
    print("v1.5 features test passed")
    print("Corrected Word:", corrected)


if __name__ == "__main__":
    run_test()
