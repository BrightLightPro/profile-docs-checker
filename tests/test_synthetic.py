from __future__ import annotations

from pathlib import Path
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import fitz
from openpyxl import Workbook, load_workbook

from doc_table_checker.core import normalize_language, preview_extraction, run_validation, summarize_results


def make_table_pdf(path: Path, rows: list[tuple[str, str]]):
    doc = fitz.open()
    doc.new_page(width=595, height=842)
    page = doc.new_page(width=595, height=842)

    x0, y0 = 70, 120
    col1_w, col2_w = 190, 250
    row_h = 28
    x1 = x0 + col1_w + col2_w
    y1 = y0 + len(rows) * row_h

    page.draw_rect(fitz.Rect(x0, y0, x1, y1), width=0.8)
    page.draw_line((x0 + col1_w, y0), (x0 + col1_w, y1), width=0.8)
    for i in range(1, len(rows)):
        y = y0 + i * row_h
        page.draw_line((x0, y), (x1, y), width=0.8)

    for i, (label, value) in enumerate(rows):
        y = y0 + i * row_h + 8
        page.insert_text((x0 + 8, y), label, fontsize=9)
        page.insert_text((x0 + col1_w + 8, y), value, fontsize=9)

    doc.save(path)
    doc.close()


def make_excel(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append(["irrelevant"])
    ws.append(["also irrelevant"])
    ws.append(["Sprache", "Dok.-Nr.", "DOK-ID", "Freigabe-/ Änd.-Nr.", "Artikel-Nr.", "Rev.", "Vers.", "Ignored"])
    ws.append(["de", "EE******", "DOC-001", "FR-1", "ART-123", "01", "02", "x"])
    ws.append(["en", "DN-200", "DOC-002", "FR-2", "ART-222", "1", "1", "x"])
    ws.append(["de", "DN-300", "DOC-DUP", "FR-3", "ART-333", "1", "1", "x"])
    ws.append(["de", "DN-301", "DOC-DUP", "FR-4", "ART-444", "2", "1", "x"])
    ws.append(["fr", "DN-400", "DOC-004", "FR-5", "ART-555", "03", "02", "x"])
    ws.append(["jp", "DN-JP", "DOC-JP", "FR-JP", "ART-JP", "A", "B", "x"])
    wb.save(path)


def make_dataset(base: Path):
    if base.exists():
        shutil.rmtree(base)
    pdf_dir = base / "pdfs"
    pdf_dir.mkdir(parents=True)
    excel = base / "master.xlsx"
    make_excel(excel)

    labels_de = [
        "Dokumentname", "Sprache", "Dokumentnummer", "DOK-ID", "Freigabe-/Änd.-Nr.",
        "Artikelnummer", "Revision", "Version", "Ausgabe"
    ]
    labels_en = [
        "Document name", "Language", "Document number", "DOC-ID", "Release/change no.",
        "Article number", "Revision", "Version", "Issue"
    ]

    def rows(labels, values):
        return list(zip(labels, values))

    make_table_pdf(pdf_dir / "01_exact_values.pdf", rows(labels_de, [
        "Ignored name", "de_DE", "EE******", "DOC-001", "FR-1", "ART-123", "01", "02", "07/2026"
    ]))
    make_table_pdf(pdf_dir / "02_mismatch.pdf", rows(labels_en, [
        "Ignored name", "en_US", "DN-200", "DOC-002", "FR-2", "WRONG-ARTICLE", "1", "1", "07.2026"
    ]))
    make_table_pdf(pdf_dir / "03_likely_wrong_dok_id.pdf", rows(labels_en, [
        "Ignored name", "fr_FR", "DN-400", "DOC-DOES-NOT-EXIST", "FR-5", "ART-555", "03", "02", "2026-07"
    ]))
    make_table_pdf(pdf_dir / "04_duplicate_resolved.pdf", rows(labels_de, [
        "Ignored name", "de_DE", "DN-300", "DOC-DUP", "FR-3", "ART-333", "1", "1", "Juli 2026"
    ]))
    make_table_pdf(pdf_dir / "05_ausgabe_mismatch.pdf", rows(labels_de, [
        "Ignored name", "de_DE", "EE******", "DOC-001", "FR-1", "ART-123", "01", "02", "08.2026"
    ]))
    make_table_pdf(pdf_dir / "06_japanese_language.pdf", rows(labels_en, [
        "Ignored name", "ja_JP", "DN-JP", "DOC-JP", "FR-JP", "ART-JP", "A", "B", "07.2026"
    ]))
    return pdf_dir, excel


def run_synthetic_test():
    base = ROOT / "test_output"
    pdf_dir, excel = make_dataset(base)
    out = base / "validation_report.xlsx"

    previews = preview_extraction(pdf_dir, page_number=2, limit=6)
    assert previews[0].values_raw["dokumentnummer"] == "EE******"
    assert previews[0].values_norm["dokumentnummer"] == "EE******"
    assert previews[0].values_norm["ausgabe"] == "2026-07"
    assert previews[5].values_norm["sprache"] == "jp"
    assert normalize_language("jp_JP") == "jp"

    results = run_validation(pdf_dir, excel, out, expected_ausgabe="07-2026", sheet_name="Master", page_number=2)
    counts = summarize_results(results)
    print("Result counts:", counts)
    assert out.exists(), "Report was not created"
    assert counts.get("OK", 0) == 2, counts
    assert counts.get("MISMATCH", 0) == 1, counts
    assert counts.get("LIKELY_WRONG_DOK_ID", 0) == 1, counts
    assert counts.get("DUPLICATE_DOK_ID_RESOLVED", 0) == 1, counts
    assert counts.get("AUSGABE_MISMATCH", 0) == 1, counts

    exact_result = next(r for r in results if r.extracted.file_name == "01_exact_values.pdf")
    doc_comp = next(c for c in exact_result.comparisons if c.field == "Dokumentnummer")
    assert doc_comp.result == "OK"
    assert doc_comp.pdf_raw == "EE******"
    assert doc_comp.reference_raw == "EE******"

    jp_result = next(r for r in results if r.extracted.file_name == "06_japanese_language.pdf")
    language_comp = next(c for c in jp_result.comparisons if c.field == "Sprache")
    assert language_comp.result == "OK"

    wb = load_workbook(out, read_only=True, data_only=True)
    assert wb.sheetnames == ["Overview", "Issues"]
    overview_headers = [c.value for c in wb["Overview"][9]]
    assert "PDF ↔ Excel" in overview_headers
    assert "Actual PDF pages" in overview_headers
    issue_headers = [c.value for c in wb["Issues"][4]]
    assert "Value A" in issue_headers and "Value B" in issue_headers
    all_issue_values = [v for row in wb["Issues"].iter_rows(values_only=True) for v in row if v is not None]
    assert "Zulassungs-/ relevant / approval relevant" not in all_issue_values
    wb.close()
    print("Synthetic test passed.")
    print("Sample report:", out)


if __name__ == "__main__":
    run_synthetic_test()
